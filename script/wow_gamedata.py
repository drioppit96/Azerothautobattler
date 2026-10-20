# -*- coding: utf-8 -*-
"""
Game-data extractor for the auto-battler project.
Writes  C:\\wow excel\\wow_gamedata.xlsx   (readable tables)
and     C:\\wow excel\\gamedata\\*.json    (machine-readable, for the game)

Sections (each one is independent; a failing section is skipped and reported):
  talents      Talents (tree, row, column, ranks, prerequisites, text per rank)
  questchains  Quest chains / prerequisites + objectives + givers (with positions)
  gathering    Herb / ore / chest nodes: skill needed, loot, where
  skinning     Skinning (and pickpocket) loot per creature
  mobspells    Spells used by creatures (spell lists, EventAI casts)
  vendors      Vendors, what they sell, prices, stock, where
  itemsets     Item sets with the set bonuses

Run:   py wow_gamedata.py            (all sections)
       py wow_gamedata.py talents vendors     (only some)
Needs wow_common.py, wow_spelldata.py, wow_dungeons.py, wow_mobs.py next to it,
and the DBC files from extract_dbc_multi.py (run it again - new tables Talent,
TalentTab, ItemSet, Lock).
"""
import json
import os
import sys
import traceback
from collections import Counter, defaultdict

from wow_common import *   # noqa
from wow_spelldata import SpellData
from wow_dungeons import LootBook
from wow_mobs import map_name

OUTPUT = os.path.join(OUT_FOLDER, "wow_gamedata.xlsx")
JSON_DIR = os.path.join(OUT_FOLDER, "gamedata")
QUALITY_NAMES = {0: "Poor", 1: "Common", 2: "Uncommon", 3: "Rare", 4: "Epic",
                 5: "Legendary", 6: "Artifact"}
LOCK_NAMES = {1: "Lockpicking", 2: "Herbalism", 3: "Mining", 4: "Disarm trap",
              5: "Open", 6: "Treasure", 7: "Gems", 8: "Close", 9: "Arm trap",
              10: "Quick open", 11: "Quick close", 12: "Tinkering", 13: "Kneeling",
              14: "Attacking", 15: "Gahz'ridian", 16: "Blasting", 17: "Slow open",
              18: "Slow close", 19: "Fishing"}


class Ctx(object):
    pass


def save_json(name, data):
    try:
        os.makedirs(JSON_DIR)
    except OSError:
        pass
    with open(os.path.join(JSON_DIR, name + ".json"), "w") as f:
        json.dump(data, f, default=str, indent=0)
    log("  wrote gamedata\\%s.json" % name)


def names_of(cur, table, key, ids, col="name"):
    out = {}
    ids = [i for i in set(ids) if i]
    for part in chunks(ids):
        cur.execute("SELECT `%s` k, `%s` n FROM `%s` WHERE `%s` IN (%s)"
                    % (key, col, table, key, ph(len(part))), part)
        for r in cur.fetchall():
            out[r["k"]] = r["n"]
    return out


def item_info(cur, ids):
    out = {}
    for part in chunks([i for i in set(ids) if i]):
        cur.execute("SELECT entry, name, Quality q, BuyPrice bp, SellPrice sp, RequiredLevel rl, "
                    "ItemLevel il, class c, subclass sc FROM item_template WHERE entry IN (%s)"
                    % ph(len(part)), part)
        for r in cur.fetchall():
            out[r["entry"]] = r
    return out


def make_resolver(cur):
    """zone resolver calibrated with quest givers (same method as wow_mobs.py)"""
    resolver = AreaResolver()
    samples = []
    try:
        cur.execute("SELECT r.id i, q.ZoneOrSort z FROM creature_questrelation r "
                    "JOIN quest_template q ON q.entry = r.quest WHERE q.ZoneOrSort > 0 LIMIT 3000")
        rows = cur.fetchall()
        sp = creature_spawns(cur, set(r["i"] for r in rows[:900]))
        seen = set()
        for r in rows:
            if r["i"] in seen or not sp.get(r["i"]):
                continue
            seen.add(r["i"])
            mp, x, y = sp[r["i"]][0]
            samples.append((mp, x, y, r["z"]))
            if len(samples) >= 300:
                break
    except Exception:
        traceback.print_exc()
    resolver.calibrate(samples)
    return resolver


def where(ctx, spawns):
    """(map, zone, subzone, x, y, all-places text) from a spawn list"""
    if not spawns:
        return "", "", "", "", "", ""
    cnt = Counter()
    step = max(1, len(spawns) // 30)
    for mp, x, y in spawns[::step]:
        z, s = ctx.resolver.locate(mp, x, y)
        cnt[(map_name(ctx.resolver, mp), z, s)] += 1
    (m_, z_, s_), _ = cnt.most_common(1)[0]
    mp, x, y = spawns[0]
    zc = Counter()
    for (m, z, s), n in cnt.items():
        zc[(z, s)] += n
    return m_, z_, s_, round(x, 1), round(y, 1), place_text(zc, 4)


def sheet(ctx, title, cols, widths, rows):
    add_plain_sheet(ctx.wb, title, cols, widths, rows)
    ctx.wb.save(OUTPUT)


# ================================================================== talents
def sec_talents(ctx):
    import wow_talents
    wow_talents.build(ctx.cur, ctx.sd, ctx.wb, save_workbook=False)
    ctx.wb.save(OUTPUT)


# ============================================================ quest chains
def sec_questchains(ctx):
    cur = ctx.cur
    qc = columns(cur, "quest_template")
    cur.execute("SELECT * FROM quest_template")
    quests = cur.fetchall()
    G = lambda r, n, d=0: g(r, qc, n, d)
    qmap = {G(r, "entry"): r for r in quests}
    parent = {}
    for qid, r in qmap.items():
        pv = abs(G(r, "prevquestid"))
        if pv and pv in qmap:
            parent[qid] = pv
    for qid, r in qmap.items():
        nx = G(r, "nextquestinchain")
        if nx and nx in qmap and nx not in parent:
            parent[nx] = qid

    def root(q):
        seen = 0
        while q in parent and seen < 60:
            q = parent[q]
            seen += 1
        return q, seen
    # relations
    giver, turnin = defaultdict(list), defaultdict(list)
    for tbl, store, kind in (("creature_questrelation", giver, "npc"),
                             ("creature_involvedrelation", turnin, "npc"),
                             ("gameobject_questrelation", giver, "object"),
                             ("gameobject_involvedrelation", turnin, "object")):
        if table_exists(cur, tbl):
            cur.execute("SELECT id, quest FROM `%s`" % tbl)
            for x in cur.fetchall():
                store[x["quest"]].append((kind, x["id"]))
    npc_ids = set(i for d in (giver, turnin) for v in d.values() for k, i in v if k == "npc")
    go_ids = set(i for d in (giver, turnin) for v in d.values() for k, i in v if k == "object")
    npc_n = names_of(cur, "creature_template", "entry", npc_ids)
    go_n = names_of(cur, "gameobject_template", "entry", go_ids)
    npc_sp = creature_spawns(cur, npc_ids)
    go_sp = go_spawns(cur, go_ids)
    loc_cache = {}

    def who(lst):
        if not lst:
            return "", "", ""
        kind, i = lst[0]
        name = (npc_n if kind == "npc" else go_n).get(i, str(i))
        key = (kind, i)
        if key not in loc_cache:
            sp = (npc_sp if kind == "npc" else go_sp).get(i, [])
            m, z, s, x, y, _ = where(ctx, sp)
            loc_cache[key] = "%s > %s (%s, %s)" % (z, s, x, y) if z else (
                "%s (%s, %s)" % (m, x, y) if m else "")
        return "%s%s" % (name, " (object)" if kind == "object" else ""), i, loc_cache[key]

    # objective names
    cids, gids, iids, sids = set(), set(), set(), set()
    for r in quests:
        for k in range(1, 5):
            v = G(r, "reqcreatureorgoid%d" % k)
            (cids if v > 0 else gids).add(abs(v)) if v else None
            iids.add(G(r, "reqitemid%d" % k))
            sids.add(G(r, "reqspellcast%d" % k))
        iids.add(G(r, "srcitemid"))
    cn = names_of(cur, "creature_template", "entry", cids)
    gn = names_of(cur, "gameobject_template", "entry", gids)
    ii = item_info(cur, iids)
    sn = spell_names(cur, sids)
    # who drops the required items
    drops = defaultdict(list)
    req_items = set(i for r in quests for k in range(1, 5) for i in [G(r, "reqitemid%d" % k)] if i)
    if req_items and table_exists(cur, "creature_loot_template"):
        le = loot_entries(cur, "creature_loot_template", sorted(req_items))
        ents = set(e for lst in le.values() for e, ref, ch in lst)
        owners = defaultdict(list)
        for part in chunks(list(ents)):
            cur.execute("SELECT Entry e, Name n, MinLevel lo, MaxLevel hi, LootId l FROM creature_template "
                        "WHERE LootId IN (%s)" % ph(len(part)), part)
            for r in cur.fetchall():
                owners[r["l"]].append("%s (%s%s)" % (r["n"], r["lo"], "" if r["lo"] == r["hi"] else "-%s" % r["hi"]))
        for item, lst in le.items():
            for e, ref, ch in lst:
                if ref:
                    continue
                for o in owners.get(e, [])[:6]:
                    drops[item].append("%s %.0f%%" % (o, ch))
    chain_rows, obj_rows, out = [], [], []
    for qid in sorted(qmap):
        r = qmap[qid]
        rt, depth = root(qid)
        gv, gid, gloc = who(giver.get(qid))
        tv, tid, tloc = who(turnin.get(qid))
        race = ", ".join(mask_names(G(r, "requiredraces"), RACES)) if G(r, "requiredraces") else "All"
        cls = ", ".join(mask_names(G(r, "requiredclasses"), CLASSES)) if G(r, "requiredclasses") else "All"
        pv = G(r, "prevquestid")
        objs = []
        for k in range(1, 5):
            v, cnt = G(r, "reqcreatureorgoid%d" % k), G(r, "reqcreatureorgocount%d" % k)
            if v:
                nm = cn.get(v, v) if v > 0 else gn.get(-v, -v)
                kind = "Kill" if v > 0 else "Use/Interact"
                objs.append(("%s %s x%d" % (kind, nm, cnt), kind, v, nm, cnt))
            it, ic = G(r, "reqitemid%d" % k), G(r, "reqitemcount%d" % k)
            if it:
                nm = ii.get(it, {}).get("name", it)
                objs.append(("Collect %s x%d" % (nm, ic), "Collect", it, nm, ic))
                obj_rows.append([qid, G(r, "title", ""), "Collect", it, nm, ic, "; ".join(drops.get(it, [])[:6])])
            sp = G(r, "reqspellcast%d" % k)
            if sp:
                objs.append(("Cast %s" % sn.get(sp, sp), "Cast", sp, sn.get(sp, ""), 1))
        for o in objs:
            if o[1] != "Collect":
                obj_rows.append([qid, G(r, "title", ""), o[1], o[2], o[3], o[4], ""])
        src = G(r, "srcitemid")
        chain_rows.append([qid, G(r, "title", ""), G(r, "questlevel"), G(r, "minlevel"),
                           G(r, "zoneorsort"), race, cls,
                           pv or "", G(r, "nextquestid") or "", G(r, "nextquestinchain") or "",
                           G(r, "exclusivegroup") or "", G(r, "breadcrumbforquestid") or "",
                           rt, qmap[rt].get(qc["title"], "") if rt != qid else "", depth + 1,
                           gv, gid, gloc, tv, tid, tloc,
                           "; ".join(o[0] for o in objs),
                           ("%s x%d" % (ii.get(src, {}).get("name", src), G(r, "srcitemcount"))) if src else "",
                           G(r, "reworreqmoney") or "",
                           "https://www.wowhead.com/classic/quest=%d" % qid])
        out.append({"id": qid, "title": G(r, "title", ""), "level": G(r, "questlevel"),
                    "min_level": G(r, "minlevel"), "zone_or_sort": G(r, "zoneorsort"),
                    "races": G(r, "requiredraces"), "classes": G(r, "requiredclasses"),
                    "prev": pv, "next": G(r, "nextquestid"), "next_in_chain": G(r, "nextquestinchain"),
                    "exclusive_group": G(r, "exclusivegroup"), "chain_root": rt, "chain_step": depth + 1,
                    "giver": [list(x) for x in giver.get(qid, [])], "turnin": [list(x) for x in turnin.get(qid, [])],
                    "objectives": [{"type": o[1], "id": o[2], "name": o[3], "count": o[4]} for o in objs],
                    "start_item": src,
                    "reward_money": G(r, "reworreqmoney"),
                    "reward_money_max_level": G(r, "rewmoneymaxlevel"),
                    "reward_items": [[G(r, "rewitemid%d" % k), G(r, "rewitemcount%d" % k)] for k in range(1, 5) if G(r, "rewitemid%d" % k)],
                    "choice_items": [[G(r, "rewchoiceitemid%d" % k), G(r, "rewchoiceitemcount%d" % k)] for k in range(1, 7) if G(r, "rewchoiceitemid%d" % k)],
                    "rep": [[G(r, "rewrepfaction%d" % k), G(r, "rewrepvalue%d" % k)] for k in range(1, 6) if G(r, "rewrepfaction%d" % k)]})
        # group quests and reputation (game/reputation.js; also written by wow_reputation.py)
        qx = out[-1]
        if G(r, "type"):
            qx["type"] = G(r, "type")
        if G(r, "suggestedplayers"):
            qx["sugg"] = G(r, "suggestedplayers")
        for key, f, v in (("minrep", "requiredminrepfaction", "requiredminrepvalue"), ("maxrep", "requiredmaxrepfaction", "requiredmaxrepvalue"),
                          ("repobj", "repobjectivefaction", "repobjectivevalue")):
            if G(r, f):
                qx[key] = [G(r, f), G(r, v)]
    sheet(ctx, "Quest Chains",
          ["Quest ID", "Title", "Level", "Min Level", "Zone/Sort ID", "Races", "Classes",
           "Requires Quest (prev)", "Next Quest", "Next In Chain", "Exclusive Group", "Breadcrumb For",
           "Chain Root ID", "Chain Root Title", "Step In Chain", "Given By", "Giver ID", "Giver Location",
           "Turn In To", "Turn-in ID", "Turn-in Location", "Objectives", "Start Item", "Money Reward",
           "Wowhead Link"],
          [8, 34, 6, 6, 9, 14, 14, 10, 9, 9, 9, 9, 9, 28, 7, 24, 8, 36, 24, 8, 36, 60, 24, 11, 34],
          chain_rows)
    sheet(ctx, "Quest Objectives",
          ["Quest ID", "Quest", "Type", "Target ID", "Target", "Count", "Dropped By (top)"],
          [8, 34, 13, 9, 34, 7, 70], obj_rows)
    save_json("quests", out)
    log("  quests: %d, objectives: %d" % (len(chain_rows), len(obj_rows)))


# =============================================================== gathering
def sec_gathering(ctx):
    cur = ctx.cur
    lock = load_dbc("Lock")
    gc = columns(cur, "gameobject_template")
    cur.execute("SELECT * FROM gameobject_template WHERE type IN (3, 25)")
    gos = cur.fetchall()
    lockinfo = {}
    prio = ["Herbalism", "Mining", "Treasure", "Lockpicking", "Open"]
    if lock:
        for r in lock.rows:
            found = []
            for k in range(8):
                if r[1 + k] == 2:       # skill-type lock entry
                    found.append((LOCK_NAMES.get(r[9 + k], "Lock %d" % r[9 + k]), r[17 + k]))
            if found:
                found.sort(key=lambda f: prio.index(f[0]) if f[0] in prio else 99)
                lockinfo[r[0]] = found[0]
    nodes = []
    for r in gos:
        lid = g(r, gc, "data0")
        kind, skill = lockinfo.get(lid, ("", 0))
        if g(r, gc, "type") == 25:
            kind, skill = "Fishing hole", 0
        if kind in ("Herbalism", "Mining", "Fishing hole", "Treasure"):
            nodes.append((r, kind, skill))
    log("  gathering node types: %d" % len(nodes))
    book = LootBook(cur, "gameobject_loot_template")
    sp = go_spawns(cur, [g(r, gc, "entry") for r, k, s in nodes])
    iids = set()
    loots = {}
    for r, kind, skill in nodes:
        lt = g(r, gc, "data1")
        loots[g(r, gc, "entry")] = book.expand(lt) if lt else {}
        iids.update(loots[g(r, gc, "entry")].keys())
    ii = item_info(cur, iids)
    rows, out = [], []
    for r, kind, skill in sorted(nodes, key=lambda n: (n[1], n[2], g(n[0], gc, "name", ""))):
        e = g(r, gc, "entry")
        m, z, s, x, y, allp = where(ctx, sp.get(e, []))
        lt = loots.get(e, {})
        txt = "; ".join("%s %.1f%% (%s-%s)" % (ii.get(i, {}).get("name", i), v[0], v[1], v[2])
                        for i, v in sorted(lt.items(), key=lambda kv: -kv[1][0])[:8])
        rows.append([e, g(r, gc, "name", ""), kind, skill or "", len(sp.get(e, [])), m, z, s, allp, txt])
        out.append({"id": e, "name": g(r, gc, "name", ""), "kind": kind, "skill": skill,
                    "loot_id": g(r, gc, "data1"),
                    "loot": [{"item": i, "chance": round(v[0], 2), "min": v[1], "max": v[2]} for i, v in lt.items()],
                    "spawns": [[a, round(b, 1), round(c, 1)] for a, b, c in sp.get(e, [])]})
    sheet(ctx, "Gathering Nodes",
          ["ID", "Node", "Type", "Skill Required", "Spawns", "Map", "Main Zone", "Main Subzone",
           "Locations", "Loot (chance, count)"],
          [8, 28, 12, 9, 8, 16, 22, 24, 50, 80], rows)
    save_json("gathering", out)


# ================================================================= skinning
def sec_skinning(ctx):
    cur = ctx.cur
    cc = columns(cur, "creature_template")
    cur.execute("SELECT * FROM creature_template WHERE SkinningLootId > 0 OR PickpocketLootId > 0")
    mobs = cur.fetchall()
    sk, pk = LootBook(cur, "skinning_loot_template"), LootBook(cur, "pickpocketing_loot_template")
    iids = set()
    data = []
    for r in mobs:
        a = sk.expand(g(r, cc, "skinninglootid")) if g(r, cc, "skinninglootid") else {}
        b = pk.expand(g(r, cc, "pickpocketlootid")) if g(r, cc, "pickpocketlootid") else {}
        iids.update(a.keys())
        iids.update(b.keys())
        data.append((r, a, b))
    ii = item_info(cur, iids)
    rows, out = [], []
    for r, a, b in sorted(data, key=lambda t: (g(t[0], cc, "minlevel"), g(t[0], cc, "name", ""))):
        lo, hi = g(r, cc, "minlevel"), g(r, cc, "maxlevel")
        need = max(1, lo * 5) if lo < 20 else lo * 5
        for kind, lt in (("Skinning", a), ("Pickpocket", b)):
            if not lt:
                continue
            txt = "; ".join("%s %.1f%% (%s-%s)" % (ii.get(i, {}).get("name", i), v[0], v[1], v[2])
                            for i, v in sorted(lt.items(), key=lambda kv: -kv[1][0])[:10])
            rows.append([g(r, cc, "entry"), g(r, cc, "name", ""), lo if lo == hi else "%d-%d" % (lo, hi),
                         FAMILY_NAMES.get(g(r, cc, "family"), ""), kind, need if kind == "Skinning" else "", txt])
            out.append({"entry": g(r, cc, "entry"), "name": g(r, cc, "name", ""), "min_level": lo, "max_level": hi,
                        "kind": kind, "loot": [{"item": i, "chance": round(v[0], 2), "min": v[1], "max": v[2]}
                                               for i, v in lt.items()]})
    sheet(ctx, "Skinning & Pickpocket", ["Mob ID", "Mob", "Level", "Family", "Kind",
                                         "Skill Needed (approx.)", "Loot"],
          [8, 30, 8, 12, 11, 10, 100], rows)
    save_json("skinning", out)


FAMILY_NAMES = {1: "Wolf", 2: "Cat", 3: "Spider", 4: "Bear", 5: "Boar", 6: "Crocolisk",
                7: "Carrion Bird", 8: "Crab", 9: "Gorilla", 11: "Raptor", 12: "Tallstrider",
                20: "Scorpid", 21: "Turtle", 24: "Bat", 25: "Hyena", 26: "Owl", 27: "Wind Serpent"}


# =============================================================== mob spells
def sec_mobspells(ctx):
    cur = ctx.cur
    cc = columns(cur, "creature_template")
    cur.execute("SELECT Entry, Name, MinLevel, MaxLevel, `Rank`, SpellList FROM creature_template")
    mobs = {r["Entry"]: r for r in cur.fetchall()}
    uses = []     # (mob, spell, source, initial_min, initial_max, repeat_min, repeat_max, chance, note)
    if table_exists(cur, "creature_spell_list"):
        lc = columns(cur, "creature_spell_list")
        lists = defaultdict(list)
        cur.execute("SELECT * FROM creature_spell_list")
        for r in cur.fetchall():
            lists[g(r, lc, "id")].append(r)
        for mid, m in mobs.items():
            for r in lists.get(m["SpellList"], []):
                uses.append((mid, g(r, lc, "spellid"), "Spell list", g(r, lc, "initialmin"), g(r, lc, "initialmax"),
                             g(r, lc, "repeatmin"), g(r, lc, "repeatmax"), g(r, lc, "probability"), ""))
    if table_exists(cur, "creature_template_spells"):
        tc = columns(cur, "creature_template_spells")
        cur.execute("SELECT * FROM creature_template_spells")
        for r in cur.fetchall():
            for k in range(1, 11):
                v = g(r, tc, "spell%d" % k)
                if v and g(r, tc, "entry") in mobs:
                    uses.append((g(r, tc, "entry"), v, "Template spells", "", "", "", "", "", ""))
    if table_exists(cur, "creature_ai_scripts"):
        ac = columns(cur, "creature_ai_scripts")
        cur.execute("SELECT * FROM creature_ai_scripts")
        EV = {0: "timer (in combat)", 1: "timer (out of combat)", 2: "HP %", 3: "mana %", 4: "aggro",
              5: "kill", 6: "death", 7: "evade", 8: "hit by spell", 9: "range", 11: "spawned",
              12: "target HP %", 13: "target casting", 14: "friendly HP", 15: "friendly CC",
              16: "missing buff", 17: "summoned", 20: "target mana %", 22: "received emote",
              23: "has aura", 24: "target has aura", 26: "LOS", 27: "missing aura"}
        for r in cur.fetchall():
            cid = g(r, ac, "creature_id")
            if cid not in mobs:
                continue
            for k in (1, 2, 3):
                if g(r, ac, "action%d_type" % k) == 11:
                    ev = g(r, ac, "event_type")
                    p = [g(r, ac, "event_param%d" % i) for i in range(1, 5)]
                    note = EV.get(ev, "event %d" % ev)
                    if ev in (0, 1):
                        uses.append((cid, g(r, ac, "action%d_param1" % k), "EventAI", p[0], p[1], p[2], p[3],
                                     g(r, ac, "event_chance"), note))
                    else:
                        uses.append((cid, g(r, ac, "action%d_param1" % k), "EventAI", "", "", "", "",
                                     g(r, ac, "event_chance"), "%s %s" % (note, [x for x in p if x])))
    sd = ctx.sd
    ids = set(u[1] for u in uses if u[1])
    info = {}
    for rr in sd.fetch(cur, ids):
        d = sd.row(rr)
        info[d["id"]] = d
    rows, out = [], []
    seen = set()
    for mid, spid, src, a, b, c, d, ch, note in sorted(uses, key=lambda u: (mobs[u[0]]["MinLevel"], mobs[u[0]]["Name"], u[1])):
        if (mid, spid, src, note) in seen or spid not in info:
            continue
        seen.add((mid, spid, src, note))
        m, sp = mobs[mid], info[spid]
        lo, hi = m["MinLevel"] or 0, m["MaxLevel"] or 0
        rows.append([mid, m["Name"], lo if lo == hi else "%d-%d" % (lo, hi),
                     {0: "Normal", 1: "Elite", 2: "Rare Elite", 3: "Boss", 4: "Rare"}.get(m["Rank"], ""),
                     spid, sp["name"], sp["rank"], sp["school"], sp["cast"], sp["cooldown"], sp["range"],
                     sp["duration"], sp["damage"], sp["effects"], src, note,
                     "%s-%s" % (a, b) if a != "" else "", "%s-%s" % (c, d) if c != "" else "", ch, sp["desc"]])
        out.append({"mob": mid, "spell": spid, "source": src, "note": note,
                    "initial_ms": [a, b], "repeat_ms": [c, d], "chance": ch})
    sheet(ctx, "Mob Spells",
          ["Mob ID", "Mob", "Level", "Rank", "Spell ID", "Spell", "Rank", "School", "Cast (s)", "Cooldown (s)",
           "Range", "Duration", "Damage / Healing", "Effects", "Source", "Trigger", "Initial Delay (ms)",
           "Repeat (ms)", "Chance %", "Description"],
          [8, 28, 8, 9, 8, 26, 8, 9, 8, 9, 9, 9, 30, 44, 14, 24, 14, 14, 8, 60], rows)
    save_json("mob_spells", {"uses": out, "spells": {str(k): v for k, v in info.items()}})
    log("  mob spell uses: %d, distinct spells: %d" % (len(rows), len(info)))


# ================================================================== vendors
def sec_vendors(ctx):
    cur = ctx.cur
    cc = columns(cur, "creature_template")
    # vanilla 1.12 flag values: vendor = 4, repair = 16384 (not 128 / 4096 as in later versions)
    nv = columns(cur, "npc_vendor")
    cur.execute("SELECT * FROM creature_template WHERE (NpcFlags & 4) > 0 "
                "OR Entry IN (SELECT DISTINCT `%s` FROM npc_vendor) "
                "OR VendorTemplateId IN (SELECT DISTINCT entry FROM npc_vendor_template)" % nv["entry"])
    vend = cur.fetchall()
    cur.execute("SELECT * FROM npc_vendor")
    direct = defaultdict(list)
    for r in cur.fetchall():
        direct[g(r, nv, "entry")].append(r)
    templ = defaultdict(list)
    tv = {}
    if table_exists(cur, "npc_vendor_template"):
        tv = columns(cur, "npc_vendor_template")
        cur.execute("SELECT * FROM npc_vendor_template")
        for r in cur.fetchall():
            templ[g(r, tv, "entry")].append(r)
    iids = set()
    for lst in direct.values():
        iids.update(g(r, nv, "item") for r in lst)
    for lst in templ.values():
        iids.update(g(r, tv, "item") for r in lst)
    ii = item_info(cur, iids)
    sp = creature_spawns(cur, [g(r, cc, "entry") for r in vend])
    rows, out = [], []
    for r in sorted(vend, key=lambda r: (g(r, cc, "name", ""), g(r, cc, "entry"))):
        e = g(r, cc, "entry")
        items = [(x, nv) for x in direct.get(e, [])]
        vt = g(r, cc, "vendortemplateid")
        if vt:
            items += [(x, tv) for x in templ.get(vt, [])]
        if not items:
            continue
        m, z, s, x_, y_, _ = where(ctx, sp.get(e, []))
        lo, hi = g(r, cc, "minlevel"), g(r, cc, "maxlevel")
        lst = []
        for x, cols_ in items:
            it = g(x, cols_, "item")
            inf = ii.get(it, {})
            mc = g(x, cols_, "maxcount")
            rows.append([e, g(r, cc, "name", ""), g(r, cc, "subname", "") or "", lo if lo == hi else "%d-%d" % (lo, hi),
                         g(r, cc, "faction"), m, z, s, x_, y_, it, inf.get("name", ""),
                         QUALITY_NAMES.get(inf.get("q"), ""), inf.get("rl", ""), money(inf.get("bp", 0)) or "",
                         inf.get("bp", 0), mc or "unlimited", g(x, cols_, "incrtime") or "",
                         g(x, cols_, "extendedcost") or ""])
            lst.append({"item": it, "price": inf.get("bp", 0), "stock": mc, "restock_s": g(x, cols_, "incrtime")})
        out.append({"entry": e, "name": g(r, cc, "name", ""), "subname": g(r, cc, "subname", ""),
                    "level": [lo, hi], "faction": g(r, cc, "faction"), "repair": bool(g(r, cc, "npcflags") & 16384),
                    "spawns": [[a, round(b, 1), round(c, 1)] for a, b, c in sp.get(e, [])[:5]], "items": lst})
    sheet(ctx, "Vendors",
          ["Vendor ID", "Vendor", "Title", "Level", "Faction ID", "Map", "Zone", "Subzone", "X", "Y",
           "Item ID", "Item", "Quality", "Req. Level", "Buy Price", "Buy Price (copper)", "Stock",
           "Restock (s)", "Extended Cost ID"],
          [9, 28, 22, 7, 8, 16, 20, 22, 8, 8, 8, 34, 10, 7, 12, 11, 10, 9, 10], rows)
    save_json("vendors", out)
    log("  vendors: %d, rows: %d" % (len(out), len(rows)))


# ================================================================= item sets
def sec_itemsets(ctx):
    cur = ctx.cur
    d = load_dbc("ItemSet")
    if not d:
        raise RuntimeError("ItemSet.dbc missing - run extract_dbc_multi.py")
    n_items = d.fc - 28            # id + 9 name + N items + 8 spells + 8 thresholds + 2 skill
    base_sp, base_th = 10 + n_items, 18 + n_items
    ic = columns(cur, "item_template")
    cur.execute("SELECT entry, name, itemset, Quality q, RequiredLevel rl, class c, subclass sc, InventoryType it "
                "FROM item_template WHERE itemset > 0")
    members = defaultdict(list)
    for r in cur.fetchall():
        members[r["itemset"]].append(r)
    sd = ctx.sd
    ids = set()
    for r in d.rows:
        ids.update(x for x in r[base_sp:base_sp + 8] if x)
    info = {}
    for rr in sd.fetch(cur, ids):
        dd = sd.row(rr)
        info[dd["id"]] = dd
    rows, out = [], []
    for r in sorted(d.rows, key=lambda r: r[0]):
        name = ""
        for i in range(1, 10):
            t = d.s(r[i])
            if t:
                name = t
                break
        if not name:
            continue
        mem = members.get(r[0], [])
        bon = [(r[base_th + k], r[base_sp + k]) for k in range(8) if r[base_sp + k]]
        rows.append([r[0], name, max([m["rl"] for m in mem] or [0]),
                     QUALITY_NAMES.get(max([m["q"] for m in mem] or [0]), ""), len(mem),
                     ", ".join(m["name"] for m in mem)] +
                    ["(%d) %s: %s" % (th, info.get(sp, {}).get("name", sp), info.get(sp, {}).get("desc", ""))
                     for th, sp in bon][:8] + [""] * (8 - min(8, len(bon))))
        out.append({"id": r[0], "name": name, "items": [m["entry"] for m in mem],
                    "bonuses": [{"pieces": th, "spell": sp, "text": info.get(sp, {}).get("desc", "")} for th, sp in bon]})
    sheet(ctx, "Item Sets",
          ["Set ID", "Set", "Req. Level", "Quality", "Pieces", "Items", "Bonus 1", "Bonus 2", "Bonus 3",
           "Bonus 4", "Bonus 5", "Bonus 6", "Bonus 7", "Bonus 8"],
          [7, 30, 8, 10, 7, 80, 50, 50, 50, 50, 50, 50, 50, 50], rows)
    save_json("itemsets", out)
    log("  item sets: %d" % len(rows))


SECTIONS = [("talents", sec_talents), ("questchains", sec_questchains), ("gathering", sec_gathering),
            ("skinning", sec_skinning), ("mobspells", sec_mobspells), ("vendors", sec_vendors),
            ("itemsets", sec_itemsets)]


def main():
    want = [a.lower() for a in sys.argv[1:]] or [n for n, f in SECTIONS]
    from openpyxl import Workbook
    import openpyxl
    ctx = Ctx()
    if os.path.exists(OUTPUT) and len(want) < len(SECTIONS):
        try:
            ctx.wb = openpyxl.load_workbook(OUTPUT)
        except Exception:
            ctx.wb = Workbook()
            ctx.wb.remove(ctx.wb.active)
    else:
        ctx.wb = Workbook()
        ctx.wb.remove(ctx.wb.active)
    log("Connecting to database ...")
    conn = connect()
    ctx.cur = conn.cursor()
    ctx.sd = SpellData(ctx.cur)
    need_zone = any(n in want for n in ("questchains", "gathering", "vendors"))
    if need_zone:
        log("Preparing zone lookup ...")
        ctx.resolver = make_resolver(ctx.cur)
    failed = []
    for name, fn in SECTIONS:
        if name not in want:
            continue
        log("== %s ==" % name)
        try:
            fn(ctx)
        except Exception as exc:
            traceback.print_exc()
            failed.append("%s (%s)" % (name, str(exc)[:80]))
    if need_zone:
        try:
            ctx.resolver.save_cache()
        except Exception:
            pass
    conn.close()
    ctx.wb.save(OUTPUT)
    log("Done. %s" % OUTPUT)
    if failed:
        log("These sections failed: " + "; ".join(failed))


if __name__ == "__main__":
    run(main)
