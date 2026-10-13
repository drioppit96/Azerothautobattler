# -*- coding: utf-8 -*-
"""
Small shared helpers for the game-data scripts (wow_loot_all, wow_hostility,
wow_races, wow_items_json, wow_taxi, wow_navgraph, wow_navgrid, wow_spells_json).
Keep it in the same folder. Nothing to run here.
"""
import json
import os

from wow_common import *   # noqa

JSON_DIR = os.path.join(OUT_FOLDER, "gamedata")


def save_json(name, data, indent=None):
    try:
        os.makedirs(JSON_DIR)
    except OSError:
        pass
    path = os.path.join(JSON_DIR, name)
    if not path.endswith(".json"):
        path += ".json"
    with open(path, "w") as f:
        json.dump(data, f, default=str, indent=indent, separators=(",", ":") if indent is None else None)
    log("  wrote %s (%d KB)" % (path, os.path.getsize(path) // 1024))
    return path


def save_text(name, text):
    try:
        os.makedirs(JSON_DIR)
    except OSError:
        pass
    path = os.path.join(JSON_DIR, name)
    with open(path, "w") as f:
        f.write(text)
    log("  wrote %s" % path)
    return path


class BigBook(object):
    """write-only workbook: fast and low memory for sheets with 100k+ rows"""

    def __init__(self, path):
        from openpyxl import Workbook
        self.path = path
        self.wb = Workbook(write_only=True)
        self.sheets = []

    def sheet(self, title, headers, widths=None):
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Font, PatternFill
        ws = self.wb.create_sheet(title[:31])
        if widths:
            for i, w in enumerate(widths, start=1):
                ws.column_dimensions[chr(64 + i) if i < 27 else "A" + chr(64 + i - 26)].width = w
        ws.freeze_panes = "A2"
        cells = []
        for h in headers:
            c = WriteOnlyCell(ws, value=h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="305496")
            cells.append(c)
        ws.append(cells)
        self.sheets.append(ws)
        return ws

    def save(self):
        self.wb.save(self.path)
        log("  wrote %s" % self.path)


def area_names():
    """area id -> (map, parent, name)"""
    out = {}
    d = load_dbc("AreaTable")
    if not d:
        return out
    idx = d.find_field(12, "Elwynn Forest")
    if idx is None:
        idx = 11
    for r in d.rows:
        out[r[0]] = (r[1], r[2], d.s(r[idx]))
    return out


def flt(u):
    import struct
    return struct.unpack("<f", struct.pack("<I", u))[0]


def name_field(d, probe_id, probe_name, default):
    idx = d.find_field(probe_id, probe_name)
    return default if idx is None else idx


def item_icons():
    """ItemDisplayInfo id -> icon name (lower case)"""
    out = {}
    d = load_dbc("ItemDisplayInfo")
    if not d:
        return out
    for r in d.rows:
        for idx in (5, 6):
            if idx < d.fc:
                t = d.s(r[idx])
                if re.match(r"^[A-Za-z0-9_\-]{3,}$", t):
                    out[r[0]] = t.lower()
                    break
    return out


def spell_icons():
    out = {}
    d = load_dbc("SpellIcon")
    if not d:
        return out
    for r in d.rows:
        t = d.s(r[1])
        if t:
            out[r[0]] = t.replace("/", "\\").split("\\")[-1].lower()
    return out


def simplify(pts, tol=2.0):
    """Douglas-Peucker on a list of (x, y, ...) points (uses x, y only), iterative"""
    n = len(pts)
    if n < 3:
        return list(pts)
    keep = [False] * n
    keep[0] = keep[n - 1] = True
    stack = [(0, n - 1)]
    while stack:
        a, b = stack.pop()
        ax, ay = pts[a][0], pts[a][1]
        bx, by = pts[b][0], pts[b][1]
        dx, dy = bx - ax, by - ay
        ll = dx * dx + dy * dy
        best, bi = -1.0, -1
        for i in range(a + 1, b):
            px, py = pts[i][0] - ax, pts[i][1] - ay
            if ll == 0:
                d = (px * px + py * py) ** 0.5
            else:
                t = max(0.0, min(1.0, (px * dx + py * dy) / ll))
                qx, qy = px - t * dx, py - t * dy
                d = (qx * qx + qy * qy) ** 0.5
            if d > best:
                best, bi = d, i
        if best > tol:
            keep[bi] = True
            stack.append((a, bi))
            stack.append((bi, b))
    return [p for p, k in zip(pts, keep) if k]
