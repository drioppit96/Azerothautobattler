# -*- coding: utf-8 -*-
"""
Shared helpers for wow_quests.py, wow_mobs.py, wow_professions.py.
Keep this file in the SAME folder as those scripts. Nothing to run here.
Works on Python 3.8 (Windows 7).
"""
import math
import os
import pickle
import re
import struct
import sys
import time
from collections import Counter, defaultdict

# =============================================================== SETTINGS
DB = dict(host="127.0.0.1", port=3310, user="root", password="123456",
          database="classicmangos")

OUT_FOLDER = r"C:\wow excel"
DBC_DIR = os.environ.get("WOW_DBC_DIR") or os.path.join(OUT_FOLDER, "dbc")      # filled by extract_dbc_multi.py
CLIENT_DATA = r"C:\GOG Games\mangos\SoloCraft\SoloCraft 1.12.1\Data"
# archives searched for map (ADT) files, newest first
ARCHIVE_ORDER = ["patch-2.mpq", "patch.mpq", "terrain.mpq", "base.mpq"]

ICON_DIR = os.path.join(OUT_FOLDER, "icons")
ICON_URL = "https://wow.zamimg.com/images/wow/icons/medium/{name}.jpg"
EMBED_ICONS = True
# ===========================================================================

try:
    os.makedirs(OUT_FOLDER)
except OSError:
    pass

ALLIANCE_MASK, HORDE_MASK = 77, 178
QUALITY = {0: "Poor", 1: "Common", 2: "Uncommon", 3: "Rare",
           4: "Epic", 5: "Legendary", 6: "Artifact"}
QUALITY_COLOR = {0: "9D9D9D", 1: "000000", 2: "1EAA00", 3: "0070DD",
                 4: "A335EE", 5: "FF8000", 6: "E6CC80"}
INV_TYPE = {0: "", 1: "Head", 2: "Neck", 3: "Shoulder", 4: "Shirt", 5: "Chest",
            6: "Waist", 7: "Legs", 8: "Feet", 9: "Wrist", 10: "Hands",
            11: "Finger", 12: "Trinket", 13: "One-Hand", 14: "Shield",
            15: "Ranged", 16: "Back", 17: "Two-Hand", 18: "Bag", 19: "Tabard",
            20: "Chest (Robe)", 21: "Main Hand", 22: "Off Hand",
            23: "Held In Off-hand", 24: "Ammo", 25: "Thrown",
            26: "Ranged", 27: "Quiver", 28: "Relic"}
STAT = {0: "Mana", 1: "Health", 3: "Agility", 4: "Strength",
        5: "Intellect", 6: "Spirit", 7: "Stamina"}
ITEM_CLASS = {0: "Consumable", 1: "Container", 2: "Weapon", 3: "Gem", 4: "Armor",
              5: "Reagent", 6: "Projectile", 7: "Trade Goods", 8: "Generic",
              9: "Recipe", 10: "Money", 11: "Quiver", 12: "Quest", 13: "Key",
              14: "Permanent", 15: "Miscellaneous"}
ITEM_SUB = {
    0: {0: "Consumable", 1: "Potion", 2: "Elixir", 3: "Flask", 4: "Scroll",
        5: "Food & Drink", 6: "Item Enhancement", 7: "Bandage", 8: "Other"},
    2: {0: "Axe", 1: "Two-Handed Axe", 2: "Bow", 3: "Gun", 4: "Mace",
        5: "Two-Handed Mace", 6: "Polearm", 7: "Sword", 8: "Two-Handed Sword",
        10: "Staff", 13: "Fist Weapon", 14: "Miscellaneous", 15: "Dagger",
        16: "Thrown", 18: "Crossbow", 19: "Wand", 20: "Fishing Pole"},
    4: {0: "Miscellaneous", 1: "Cloth", 2: "Leather", 3: "Mail", 4: "Plate",
        6: "Shield", 7: "Libram", 8: "Idol", 9: "Totem"},
    7: {0: "Trade Goods", 1: "Parts", 2: "Explosives", 3: "Devices",
        4: "Jewelcrafting", 5: "Cloth", 6: "Leather", 7: "Metal & Stone",
        8: "Meat", 9: "Herb", 10: "Elemental", 11: "Other", 12: "Enchanting"},
    9: {0: "Book", 1: "Leatherworking", 2: "Tailoring", 3: "Engineering",
        4: "Blacksmithing", 5: "Cooking", 6: "Alchemy", 7: "First Aid",
        8: "Enchanting", 9: "Fishing"},
}
BONDING = {0: "", 1: "Binds when picked up", 2: "Binds when equipped",
           3: "Binds when used", 4: "Quest item"}
SPELL_TRIGGER = {0: "Use", 1: "Equip", 2: "Chance on hit", 4: "Soulstone",
                 5: "Use", 6: "Learn"}
RACES = [(1, "Human"), (2, "Orc"), (4, "Dwarf"), (8, "Night Elf"),
         (16, "Undead"), (32, "Tauren"), (64, "Gnome"), (128, "Troll")]
CLASSES = [(1, "Warrior"), (2, "Paladin"), (4, "Hunter"), (8, "Rogue"),
           (16, "Priest"), (64, "Shaman"), (128, "Mage"), (256, "Warlock"),
           (1024, "Druid")]
MAP_NAMES = {0: "Eastern Kingdoms", 1: "Kalimdor"}


def log(msg):
    print(str(msg).encode("ascii", "replace").decode("ascii"))
    sys.stdout.flush()


# ------------------------------------------------------------ DB helpers
def connect():
    import pymysql
    from pymysql.cursors import DictCursor
    return pymysql.connect(cursorclass=DictCursor, charset="utf8", **DB)


def table_exists(cur, name):
    cur.execute("SHOW TABLES LIKE %s", (name,))
    return cur.fetchone() is not None


def columns(cur, table):
    cur.execute("SHOW COLUMNS FROM `%s`" % table)
    return {r["Field"].lower(): r["Field"] for r in cur.fetchall()}


def chunks(seq, n=400):
    seq = list(seq)
    for i in range(0, len(seq), n):
        yield seq[i:i + n]


def ph(n):
    return ",".join(["%s"] * n)


def g(row, cols, name, default=0):
    key = cols.get(name.lower())
    if key is None:
        return default
    v = row.get(key)
    return default if v is None else v


def money(copper):
    copper = int(copper or 0)
    if copper == 0:
        return ""
    gold, rest = divmod(copper, 10000)
    silver, cop = divmod(rest, 100)
    parts = []
    if gold:
        parts.append("%dg" % gold)
    if silver:
        parts.append("%ds" % silver)
    if cop or not parts:
        parts.append("%dc" % cop)
    return " ".join(parts)


def clean_text(t):
    """Turn quest text variables like $N, $B, $g...; into readable text."""
    if not t:
        return ""
    t = re.sub(r"\$[gG]([^:;]*):([^;]*);", r"\1/\2", t)
    t = t.replace("$B", " ").replace("$b", " ")
    t = t.replace("$N", "<name>").replace("$n", "<name>")
    t = t.replace("$C", "<class>").replace("$c", "<class>")
    t = t.replace("$R", "<race>").replace("$r", "<race>")
    t = t.replace("\r", " ").replace("\n", " ")
    return re.sub(r"\s+", " ", t).strip()


def mask_names(mask, table):
    mask = int(mask or 0)
    return [n for bit, n in table if mask & bit]


def side_from_races(mask):
    mask = int(mask or 0)
    if mask in (0, -1) or (mask & 255) == 255:
        return ""
    a = bool(mask & ALLIANCE_MASK)
    h = bool(mask & HORDE_MASK)
    if a and not h:
        return "Alliance"
    if h and not a:
        return "Horde"
    return ""


# ------------------------------------------------------------------- DBC
class DBC(object):
    def __init__(self, path):
        with open(path, "rb") as f:
            data = f.read()
        magic, rc, fc, rs, ss = struct.unpack("<4sIIII", data[:20])
        if magic != b"WDBC":
            raise ValueError("not a DBC file: %s" % path)
        self.rc, self.fc, self.rs = rc, fc, rs
        self.strings = data[20 + rc * rs:20 + rc * rs + ss]
        fmt = "<%dI" % fc
        self.rows = [struct.unpack_from(fmt, data, 20 + i * rs) for i in range(rc)]

    def s(self, off):
        st = self.strings
        if off <= 0 or off >= len(st) or st[off - 1] != 0:
            return ""
        end = st.find(b"\0", off)
        return st[off:end].decode("utf-8", "ignore")

    def find_field(self, row_id, expected):
        """index of the string field whose value for row_id equals expected"""
        for r in self.rows:
            if r[0] == row_id:
                for i in range(1, self.fc):
                    if self.s(r[i]) == expected:
                        return i
        return None


def load_dbc(name):
    path = os.path.join(DBC_DIR, name + ".dbc")
    if not os.path.exists(path):
        log("  (missing %s - run extract_dbc_multi.py)" % path)
        return None
    try:
        return DBC(path)
    except Exception as exc:
        log("  could not read %s: %s" % (path, exc))
        return None


def load_faction_names():
    d = load_dbc("Faction")
    if not d:
        return {}
    idx = d.find_field(72, "Stormwind")
    if idx is None:
        idx = 20
    return {r[0]: d.s(r[idx]) for r in d.rows}


def load_faction_groups():
    """FactionTemplate id -> group bitmask (2 = Alliance, 4 = Horde)"""
    d = load_dbc("FactionTemplate")
    if not d:
        return {}
    return {r[0]: r[3] for r in d.rows}


def side_from_template(faction_id, groups):
    grp = groups.get(faction_id, 0)
    a, h = bool(grp & 2), bool(grp & 4)
    if a and not h:
        return "Alliance"
    if h and not a:
        return "Horde"
    return ""


# ----------------------------------------------------------- zone lookup
class AreaResolver(object):
    """(map, x, y) -> (zone, subzone) using the terrain files in the client."""
    TILE = 533.33333

    def __init__(self):
        self.areas = {}        # id -> (map, parent, name)
        self.map_dirs = {}
        self.inst_zone = {}
        self.tiles = {}
        self.archives = None
        self.hyp = (False, False)
        self.enabled = False
        self.cache_path = os.path.join(OUT_FOLDER, "adt_area_cache.pkl")
        self.dirty = False
        d = load_dbc("AreaTable")
        if d:
            idx = d.find_field(12, "Elwynn Forest")
            if idx is None:
                idx = 11
            for r in d.rows:
                self.areas[r[0]] = (r[1], r[2], d.s(r[idx]))
            for aid, (mp, parent, name) in sorted(self.areas.items()):
                if parent == 0 and mp not in (0, 1) and mp not in self.inst_zone:
                    self.inst_zone[mp] = aid
            log("Area names loaded: %d" % len(self.areas))
        m = load_dbc("Map")
        if m:
            for r in m.rows:
                self.map_dirs[r[0]] = m.s(r[1])
        if self.areas and self.map_dirs:
            try:
                import mpyq  # noqa: F401
                self.enabled = True
            except ImportError:
                log("  'mpyq' missing -> subzones unavailable (py -m pip install mpyq)")
        if os.path.exists(self.cache_path):
            try:
                with open(self.cache_path, "rb") as f:
                    self.tiles = pickle.load(f)
            except Exception:
                self.tiles = {}

    # -- names
    def zone_id(self, area):
        a = self.areas.get(area)
        if not a:
            return 0
        return a[1] if a[1] in self.areas else area

    def zone_name(self, zone_id):
        a = self.areas.get(zone_id)
        return a[2] if a else ""

    def names(self, area):
        a = self.areas.get(area)
        if not a:
            return ("", "")
        if a[1] and a[1] in self.areas:
            return (self.areas[a[1]][2], a[2])
        return (a[2], "")

    # -- tiles
    def _open(self):
        from mpyq import MPQArchive
        self.archives = []
        if not os.path.isdir(CLIENT_DATA):
            log("  client Data folder not found: %s" % CLIENT_DATA)
            return
        files = {f.lower(): f for f in os.listdir(CLIENT_DATA)}
        order = sorted([f for f in files if f.startswith("patch") and f.endswith(".mpq")],
                       reverse=True)
        order += [n for n in ARCHIVE_ORDER if n not in order]
        for n in order:
            if n in files:
                try:
                    self.archives.append(MPQArchive(os.path.join(CLIENT_DATA, files[n]),
                                                    listfile=False))
                except Exception as exc:
                    log("  could not open %s: %s" % (n, str(exc)[:80]))

    @staticmethod
    def parse_adt(data):
        out = {}
        pos, n = 0, len(data)
        while pos + 8 <= n:
            cid = data[pos:pos + 4]
            size = struct.unpack_from("<I", data, pos + 4)[0]
            if cid == b"KNCM" and pos + 8 + 56 <= n:
                ix, iy = struct.unpack_from("<II", data, pos + 8 + 4)
                area = struct.unpack_from("<I", data, pos + 8 + 52)[0]
                out[(ix, iy)] = area
            pos += 8 + size
        return out

    def tile(self, mp, tx, ty):
        key = (mp, tx, ty)
        if key in self.tiles:
            return self.tiles[key]
        res = None
        d = self.map_dirs.get(mp)
        if d and 0 <= tx < 64 and 0 <= ty < 64:
            if self.archives is None:
                self._open()
            fn = "World\\Maps\\%s\\%s_%d_%d.adt" % (d, d, tx, ty)
            for arc in self.archives:
                try:
                    data = arc.read_file(fn)
                except Exception:
                    data = None
                if data:
                    try:
                        res = self.parse_adt(data)
                    except Exception:
                        res = None
                    break
        self.tiles[key] = res
        self.dirty = True
        return res

    def report(self):
        found = sum(1 for v in self.tiles.values() if v)
        missing = sum(1 for v in self.tiles.values() if not v)
        log("Map tiles read: %d found, %d not found" % (found, missing))
        if self.enabled and not found:
            log("  WARNING: no terrain tiles could be read from the client, so zones")
            log("  will be blank. Check CLIENT_DATA in wow_common.py and that mpyq works.")

    def save_cache(self):
        if not self.dirty:
            return
        try:
            with open(self.cache_path, "wb") as f:
                pickle.dump(self.tiles, f, protocol=2)
        except Exception:
            pass

    def area_at(self, mp, x, y, hyp=None):
        if not self.enabled:
            return 0
        swap_tile, swap_chunk = hyp if hyp is not None else self.hyp
        a = 32.0 - float(y) / self.TILE
        b = 32.0 - float(x) / self.TILE
        f1, f2 = (b, a) if swap_tile else (a, b)
        tx, ty = int(math.floor(f1)), int(math.floor(f2))
        u, v = f1 - tx, f2 - ty
        cx, cy = int(u * 16), int(v * 16)
        if swap_chunk:
            cx, cy = cy, cx
        t = self.tile(mp, tx, ty)
        if not t:
            return 0
        return t.get((cx, cy), 0)

    def calibrate(self, samples):
        """samples: [(map, x, y, expected_zone_id)] -> picks the axis layout
        that matches the known zones best."""
        if not self.enabled or not samples:
            return
        best, best_hits = self.hyp, -1
        for hyp in ((False, False), (True, True), (False, True), (True, False)):
            hits = 0
            for mp, x, y, zid in samples:
                if self.zone_id(self.area_at(mp, x, y, hyp)) == zid:
                    hits += 1
            log("  map layout %s: %d / %d spawns match their quest zone"
                % (hyp, hits, len(samples)))
            if hits > best_hits:
                best, best_hits = hyp, hits
        self.hyp = best
        if best_hits <= len(samples) * 0.2:
            log("  WARNING: zone lookup looks unreliable (few matches).")

    def locate(self, mp, x, y):
        area = self.area_at(mp, x, y)
        if area:
            return self.names(area)
        if mp in self.inst_zone:
            return (self.areas[self.inst_zone[mp]][2], "")
        return ("", "")


# ------------------------------------------------------------ spawn data
def _spawn_rows(cur, sql, args):
    cur.execute(sql, args)
    return cur.fetchall()


def creature_spawns(cur, entries=None):
    """entry -> [(map, x, y)]  (all entries when entries is None)"""
    out = defaultdict(list)
    has_se = table_exists(cur, "creature_spawn_entry")
    parts = [None] if entries is None else list(chunks(list(entries)))
    for part in parts:
        if part is None:
            rows = _spawn_rows(cur, "SELECT id i, map m, position_x x, position_y y "
                                    "FROM creature WHERE id <> 0", ())
        else:
            rows = _spawn_rows(cur, "SELECT id i, map m, position_x x, position_y y "
                                    "FROM creature WHERE id IN (%s)" % ph(len(part)), part)
        for r in rows:
            out[r["i"]].append((r["m"], float(r["x"]), float(r["y"])))
        if has_se:
            if part is None:
                rows = _spawn_rows(cur, "SELECT s.entry i, c.map m, c.position_x x, c.position_y y "
                                        "FROM creature_spawn_entry s JOIN creature c ON c.guid = s.guid", ())
            else:
                rows = _spawn_rows(cur, "SELECT s.entry i, c.map m, c.position_x x, c.position_y y "
                                        "FROM creature_spawn_entry s JOIN creature c ON c.guid = s.guid "
                                        "WHERE s.entry IN (%s)" % ph(len(part)), part)
            for r in rows:
                out[r["i"]].append((r["m"], float(r["x"]), float(r["y"])))
    return out


def go_spawns(cur, entries):
    out = defaultdict(list)
    has_se = table_exists(cur, "gameobject_spawn_entry")
    for part in chunks(list(entries)):
        for r in _spawn_rows(cur, "SELECT id i, map m, position_x x, position_y y "
                                  "FROM gameobject WHERE id IN (%s)" % ph(len(part)), part):
            out[r["i"]].append((r["m"], float(r["x"]), float(r["y"])))
        if has_se:
            for r in _spawn_rows(cur, "SELECT s.entry i, c.map m, c.position_x x, c.position_y y "
                                      "FROM gameobject_spawn_entry s JOIN gameobject c ON c.guid = s.guid "
                                      "WHERE s.entry IN (%s)" % ph(len(part)), part):
                out[r["i"]].append((r["m"], float(r["x"]), float(r["y"])))
    return out


def place_counter(resolver, spawns, sample=40):
    """Counter {(zone, subzone): n} from (up to `sample` of) the spawns"""
    cnt = Counter()
    if not spawns:
        return cnt
    step = max(1, len(spawns) // sample) if sample else 1
    for mp, x, y in spawns[::step]:
        cnt[resolver.locate(mp, x, y)] += 1
    return cnt


def place_text(cnt, maxn=3):
    parts = []
    for (zone, sub), n in cnt.most_common(maxn):
        if not zone and not sub:
            continue
        parts.append("%s > %s" % (zone, sub) if sub else zone)
    extra = len([k for k in cnt if k != ("", "")]) - len(parts)
    txt = "; ".join(parts)
    if extra > 0:
        txt += " (+%d more)" % extra
    return txt


def main_place(cnt):
    for (zone, sub), n in cnt.most_common():
        if zone or sub:
            return zone, sub
    return "", ""


# ------------------------------------------------------------- item data
def type_text(cls, sub):
    cname = ITEM_CLASS.get(cls, "Class %s" % cls)
    sname = ITEM_SUB.get(cls, {}).get(sub)
    if sname is None:
        return cname
    if cls in (2, 4, 7, 9, 0) and sname != cname:
        return "%s - %s" % (cname, sname)
    return cname


_spell_names = {}


def spell_names(cur, ids):
    need = [i for i in set(ids) if i and i not in _spell_names]
    for part in chunks(need):
        cur.execute("SELECT Id i, SpellName n FROM spell_template WHERE Id IN (%s)"
                    % ph(len(part)), part)
        for r in cur.fetchall():
            _spell_names[r["i"]] = r["n"]
    return _spell_names


def load_items_by_ids(cur, ids):
    ic = columns(cur, "item_template")
    ids = sorted(set(int(i) for i in ids if i))
    rows = []
    for part in chunks(ids):
        cur.execute("SELECT * FROM item_template WHERE `%s` IN (%s)"
                    % (ic["entry"], ph(len(part))), part)
        rows.extend(cur.fetchall())
    spell_ids = []
    for r in rows:
        for n in range(1, 6):
            spell_ids.append(g(r, ic, "spellid_%d" % n))
    spell_names(cur, spell_ids)
    skill_names = {}
    sk = load_dbc("SkillLine")
    if sk:
        idx = sk.find_field(171, "Alchemy")
        if idx is not None:
            skill_names = {r[0]: sk.s(r[idx]) for r in sk.rows}

    items = {}
    for r in rows:
        entry = int(g(r, ic, "entry"))
        stats = []
        for n in range(1, 11):
            st, val = g(r, ic, "stat_type%d" % n), g(r, ic, "stat_value%d" % n)
            if val:
                stats.append("%+d %s" % (val, STAT.get(st, "Stat %s" % st)))
        res = []
        for key, label in (("holy_res", "Holy"), ("fire_res", "Fire"),
                           ("nature_res", "Nature"), ("frost_res", "Frost"),
                           ("shadow_res", "Shadow"), ("arcane_res", "Arcane")):
            v = g(r, ic, key)
            if v:
                res.append("%+d %s" % (v, label))
        dmin, dmax = g(r, ic, "dmg_min1", 0), g(r, ic, "dmg_max1", 0)
        delay = g(r, ic, "delay", 0)
        damage = speed = dps = ""
        if dmax and delay:
            damage = "%d - %d" % (dmin, dmax)
            speed = round(delay / 1000.0, 2)
            dps = round((dmin + dmax) / 2.0 / (delay / 1000.0), 1)
        eff = []
        for n in range(1, 6):
            sid = g(r, ic, "spellid_%d" % n)
            if sid:
                trig = SPELL_TRIGGER.get(g(r, ic, "spelltrigger_%d" % n), "Effect")
                eff.append("%s: %s" % (trig, _spell_names.get(sid, "spell %d" % sid)))
        cls, sub = g(r, ic, "class"), g(r, ic, "subclass")
        cmask = int(g(r, ic, "allowableclass", -1))
        classes = ""
        if cmask not in (0, -1) and (cmask & 1535) != 1535:
            classes = ", ".join(mask_names(cmask, CLASSES))
        rs = g(r, ic, "requiredskill")
        items[entry] = {
            "id": entry,
            "name": g(r, ic, "name", ""),
            "display": g(r, ic, "displayid"),
            "quality": g(r, ic, "quality"),
            "ilvl": g(r, ic, "itemlevel"),
            "reqlvl": g(r, ic, "requiredlevel"),
            "cls": cls, "sub": sub,
            "type": type_text(cls, sub),
            "slot": INV_TYPE.get(g(r, ic, "inventorytype"), ""),
            "binds": BONDING.get(g(r, ic, "bonding"), ""),
            "armor": g(r, ic, "armor") or "",
            "block": g(r, ic, "block") or "",
            "damage": damage, "speed": speed, "dps": dps,
            "stats": ", ".join(stats),
            "res": ", ".join(res),
            "effects": "; ".join(eff),
            "side": side_from_races(g(r, ic, "allowablerace", -1)),
            "classes": classes,
            "reqskill": ("%s %d" % (skill_names.get(rs, "skill %d" % rs),
                                    g(r, ic, "requiredskillrank"))) if rs else "",
            "stack": g(r, ic, "stackable") if g(r, ic, "stackable") > 1 else "",
            "buy": g(r, ic, "buyprice"),
            "sell": g(r, ic, "sellprice"),
            "desc": g(r, ic, "description", "") or "",
        }
    return items


ITEM_COLUMNS = [
    ("ID", 8, "id"), ("Name", 36, "name"), ("Quality", 11, "quality"),
    ("Item Level", 10, "ilvl"), ("Req Level", 10, "reqlvl"),
    ("Type", 24, "type"), ("Slot", 14, "slot"), ("Binds", 18, "binds"),
    ("Armor", 8, "armor"), ("Block", 8, "block"), ("Damage", 12, "damage"),
    ("Speed", 8, "speed"), ("DPS", 8, "dps"), ("Stats", 40, "stats"),
    ("Resistances", 20, "res"), ("Item Effects", 44, "effects"),
    ("Side", 10, "side"), ("Classes", 18, "classes"),
    ("Req Skill", 20, "reqskill"), ("Stack", 7, "stack"),
    ("Buy Price", 11, "buy"), ("Sell Price", 11, "sell"),
    ("Description", 40, "desc"),
]


def item_values(it):
    out = []
    for _, _, key in ITEM_COLUMNS:
        v = it.get(key, "")
        if key == "quality":
            v = QUALITY.get(v, "")
        elif key in ("buy", "sell"):
            v = money(v)
        out.append(v)
    return out


# --------------------------------------------------------------- sources
def loot_entries(cur, table, item_ids):
    """item -> list of (loot_entry, via_reference, chance)"""
    if not table_exists(cur, table):
        return {}
    c = columns(cur, table)
    e, i = c["entry"], c["item"]
    ch = c.get("chanceorquestchance") or c.get("chance")
    rf = c.get("mincountorref")
    out = defaultdict(list)
    for part in chunks(item_ids):
        cur.execute("SELECT `%s` e, `%s` i, %s c FROM `%s` WHERE `%s` IN (%s)"
                    % (e, i, ("`%s`" % ch) if ch else "0", table, i,
                       ph(len(part))), part)
        for r in cur.fetchall():
            out[r["i"]].append((r["e"], False, abs(float(r["c"] or 0))))
    if rf and table_exists(cur, "reference_loot_template"):
        rc = columns(cur, "reference_loot_template")
        ref_items = defaultdict(set)
        for part in chunks(item_ids):
            cur.execute("SELECT `%s` e, `%s` i FROM reference_loot_template "
                        "WHERE `%s` IN (%s)" % (rc["entry"], rc["item"],
                                                rc["item"], ph(len(part))), part)
            for r in cur.fetchall():
                ref_items[r["e"]].add(r["i"])
        if ref_items:
            for part in chunks(list(ref_items)):
                cur.execute("SELECT `%s` e, `%s` r FROM `%s` WHERE `%s` IN (%s)"
                            % (e, rf, table, rf, ph(len(part))), [-x for x in part])
                for r in cur.fetchall():
                    for it in ref_items.get(-r["r"], ()):
                        out[it].append((r["e"], True, 0.0))
    return out


def build_sources(cur, items):
    """list of (item_id, type, source name, detail) - drops, chests, vendors, quests"""
    ids = list(items)
    src = []
    if table_exists(cur, "creature_template"):
        ct = columns(cur, "creature_template")
        loot_col = ct.get("lootid")
        lv_min, lv_max = ct.get("minlevel"), ct.get("maxlevel")
        drops = loot_entries(cur, "creature_loot_template", ids)
        entries = {e for lst in drops.values() for e, _, _ in lst}
        names = defaultdict(list)
        if loot_col and entries:
            for part in chunks(list(entries)):
                cur.execute("SELECT `%s` n, `%s` l, %s lo, %s hi FROM creature_template "
                            "WHERE `%s` IN (%s)"
                            % (ct["name"], loot_col,
                               ("`%s`" % lv_min) if lv_min else "0",
                               ("`%s`" % lv_max) if lv_max else "0",
                               loot_col, ph(len(part))), part)
                for r in cur.fetchall():
                    names[r["l"]].append((r["n"], r["lo"], r["hi"]))
        for it, lst in drops.items():
            seen = set()
            for e, via_ref, chance in lst:
                if (e, via_ref) in seen:
                    continue
                seen.add((e, via_ref))
                cr = names.get(e, [])
                if cr:
                    first = ", ".join(sorted({n for n, _, _ in cr})[:3])
                    extra = len({n for n, _, _ in cr}) - 3
                    label = first + (" (+%d more)" % extra if extra > 0 else "")
                    lows = [lo for _, lo, _ in cr if lo]
                    highs = [hi for _, _, hi in cr if hi]
                    lvl = (" | creature lvl %s-%s" % (min(lows), max(highs))
                           if lows and highs else "")
                else:
                    label = "Creature loot table #%s" % e
                    lvl = ""
                detail = ("via shared loot table" if via_ref
                          else ("%.1f%% drop chance" % chance if chance else "drop"))
                src.append((it, "Drop", label, detail + lvl))

    if table_exists(cur, "gameobject_template"):
        gt = columns(cur, "gameobject_template")
        gdrops = loot_entries(cur, "gameobject_loot_template", ids)
        entries = {e for lst in gdrops.values() for e, _, _ in lst}
        gnames = defaultdict(list)
        if entries and "data1" in gt:
            for part in chunks(list(entries)):
                cur.execute("SELECT `%s` n, `%s` l FROM gameobject_template "
                            "WHERE `%s` IN (3, 25) AND `%s` IN (%s)"
                            % (gt["name"], gt["data1"], gt["type"], gt["data1"],
                               ph(len(part))), part)
                for r in cur.fetchall():
                    gnames[r["l"]].append(r["n"])
        for it, lst in gdrops.items():
            for e, via_ref, chance in lst:
                nm = gnames.get(e)
                if not nm:
                    continue
                label = ", ".join(sorted(set(nm))[:3])
                detail = ("via shared loot table" if via_ref
                          else ("%.1f%% chance" % chance if chance else "loot"))
                src.append((it, "Chest/Object", label, detail))

    for vt, tpl_col in (("npc_vendor", None), ("npc_vendor_template", "vendortemplateid")):
        if not table_exists(cur, vt) or not table_exists(cur, "creature_template"):
            continue
        vc = columns(cur, vt)
        ct = columns(cur, "creature_template")
        key = ct["entry"] if tpl_col is None else ct.get(tpl_col)
        if not key:
            continue
        ext = vc.get("extendedcost")
        for part in chunks(ids):
            cur.execute("SELECT v.`%s` e, v.`%s` i, %s x, c.`%s` n FROM `%s` v "
                        "JOIN creature_template c ON c.`%s` = v.`%s` "
                        "WHERE v.`%s` IN (%s)"
                        % (vc["entry"], vc["item"],
                           ("v.`%s`" % ext) if ext else "0",
                           ct["name"], vt, key, vc["entry"], vc["item"],
                           ph(len(part))), part)
            seen = set()
            for r in cur.fetchall():
                k = (r["i"], r["n"])
                if k in seen:
                    continue
                seen.add(k)
                price = ("Special cost (honor/tokens)" if r["x"]
                         else money(items[r["i"]]["buy"]) or "free")
                src.append((r["i"], "Vendor", r["n"], price))

    if table_exists(cur, "quest_template"):
        qc = columns(cur, "quest_template")
        rew = [(k, v) for k, v in qc.items() if re.match(r"rew(choice)?itemid\d+$", k)]
        if rew:
            idset = set(ids)
            sel = ", ".join("`%s`" % v for _, v in rew)
            cur.execute("SELECT `%s` e, `%s` t, %s FROM quest_template"
                        % (qc["entry"], qc.get("title", qc["entry"]), sel))
            for r in cur.fetchall():
                for k, v in rew:
                    item = r.get(v)
                    if item in idset:
                        kind = "choice reward" if "choice" in k else "reward"
                        src.append((item, "Quest", r["t"], kind))
    return src


def source_summary(srcs):
    counts = defaultdict(int)
    for typ, _, _ in srcs:
        counts[typ] += 1
    return ", ".join("%s (%d)" % (k, v) for k, v in sorted(counts.items()))


# ------------------------------------------------------------------ icons
_ICON_FAILS = set()


def get_icon(name):
    if not name:
        return None
    import requests
    try:
        os.makedirs(ICON_DIR)
    except OSError:
        pass
    for ext in (".png", ".jpg"):
        cached = os.path.join(ICON_DIR, name.lower() + ext)
        if os.path.exists(cached):
            return cached
    if name.lower() in _ICON_FAILS:
        return None
    path = os.path.join(ICON_DIR, name.lower() + ".jpg")
    try:
        r = requests.get(ICON_URL.format(name=name.lower()), timeout=20,
                         headers={"User-Agent": "Mozilla/5.0"})
        if r.status_code == 200 and r.content:
            with open(path, "wb") as f:
                f.write(r.content)
            time.sleep(0.05)
            return path
    except Exception:
        pass
    _ICON_FAILS.add(name.lower())
    return None


class Icons(object):
    """item display id / spell icon id -> icon file path"""

    def __init__(self):
        self.item_names = {}
        self.spell_names = {}
        self.missing = 0
        if not EMBED_ICONS:
            return
        d = load_dbc("ItemDisplayInfo")
        if d:
            for r in d.rows:
                for idx in (5, 6):
                    if idx < d.fc:
                        t = d.s(r[idx])
                        if re.match(r"^[A-Za-z0-9_\-]{3,}$", t):
                            self.item_names[r[0]] = t
                            break
            log("Item icon names loaded: %d" % len(self.item_names))
        d = load_dbc("SpellIcon")
        if d:
            for r in d.rows:
                t = d.s(r[1])
                if t:
                    self.spell_names[r[0]] = t.replace("/", "\\").split("\\")[-1]
            log("Spell icon names loaded: %d" % len(self.spell_names))

    def for_item(self, display_id):
        if not EMBED_ICONS:
            return None
        p = get_icon(self.item_names.get(display_id, ""))
        if not p:
            self.missing += 1
        return p

    def for_spell(self, icon_id):
        if not EMBED_ICONS:
            return None
        p = get_icon(self.spell_names.get(icon_id, ""))
        if not p:
            self.missing += 1
        return p


# ------------------------------------------------------------------ excel
def _openpyxl():
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    return XLImage, Alignment, Font, PatternFill, get_column_letter


def new_sheet(wb, title, cols, widths, first=False):
    XLImage, Alignment, Font, PatternFill, gcl = _openpyxl()
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(cols)
    fill = PatternFill("solid", fgColor="1F3864")
    for i, w in enumerate(widths, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[gcl(i)].width = w
    return ws


def finish_sheet(ws, ncols, freeze="A2"):
    XLImage, Alignment, Font, PatternFill, gcl = _openpyxl()
    ws.freeze_panes = freeze
    ws.auto_filter.ref = "A1:%s%d" % (gcl(ncols), max(ws.max_row, 2))


def style_row(ws, row, ncols, link_col=None, link=None, name_col=None, quality=None,
              icon_path=None, icon_col=1, tall=False):
    XLImage, Alignment, Font, PatternFill, gcl = _openpyxl()
    if tall:
        ws.row_dimensions[row].height = 30
    for col in range(1, ncols + 1):
        ws.cell(row=row, column=col).alignment = Alignment(vertical="center")
    if name_col and quality is not None:
        ws.cell(row=row, column=name_col).font = Font(
            bold=True, color=QUALITY_COLOR.get(quality, "000000"))
    if link_col and link:
        c = ws.cell(row=row, column=link_col)
        c.hyperlink = link
        c.font = Font(color="0563C1", underline="single")
    if icon_path:
        try:
            ws.add_image(XLImage(icon_path), "%s%d" % (gcl(icon_col), row))
        except Exception:
            pass


def write_item_sheet(wb, title, items, icons, extra_headers=(), extra_widths=(),
                     extra_fn=None, first=False, before=(), before_widths=(),
                     before_fn=None):
    """One row per entry of `items` (list of item dicts, optionally with extra
    data via before_fn / extra_fn). Column A holds the icon."""
    cols = (["Icon"] + list(before) + [c[0] for c in ITEM_COLUMNS] + ["Wowhead Link"]
            + list(extra_headers))
    widths = ([8] + list(before_widths) + [c[1] for c in ITEM_COLUMNS] + [34]
              + list(extra_widths))
    ws = new_sheet(wb, title, cols, widths, first=first)
    nb = len(before)
    name_col = 2 + nb + 1
    link_col = 1 + nb + len(ITEM_COLUMNS) + 1
    for n, it in enumerate(items, start=2):
        link = "https://www.wowhead.com/classic/item=%d" % it["id"]
        row = ([""] + (list(before_fn(it)) if before_fn else []) + item_values(it)
               + [link] + (list(extra_fn(it)) if extra_fn else []))
        ws.append(row)
        style_row(ws, n, len(cols), link_col=link_col, link=link, name_col=name_col,
                  quality=it["quality"], icon_path=icons.for_item(it["display"]),
                  tall=True)
        if (n - 1) % 100 == 0:
            log("  %s: %d / %d rows" % (title, n - 1, len(items)))
    finish_sheet(ws, len(cols), freeze="A2")
    return ws


def add_plain_sheet(wb, title, cols, widths, rows, first=False, link_col=None):
    ws = new_sheet(wb, title, cols, widths, first=first)
    XLImage, Alignment, Font, PatternFill, gcl = _openpyxl()
    for n, row in enumerate(rows, start=2):
        ws.append(row)
        if link_col and row[link_col - 1]:
            c = ws.cell(row=n, column=link_col)
            c.hyperlink = row[link_col - 1]
            c.font = Font(color="0563C1", underline="single")
    finish_sheet(ws, len(cols))
    return ws


def run(main):
    import traceback
    try:
        main()
    except Exception:
        traceback.print_exc()
    input("\nPress Enter to close this window...")
