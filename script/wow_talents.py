# -*- coding: utf-8 -*-
"""
Talents, talent trees and their in-game layout.

Writes
  C:\\wow excel\\wow_talents.xlsx          sheets: Talent Trees, Talents (every rank), Talent Layout (grid like the game)
  C:\\wow excel\\gamedata\\talents.json     machine-readable (for the game)
  C:\\wow excel\\gamedata\\talent_trees.html   open in a browser: all 27 trees, clickable calculator
  C:\\wow excel\\gamedata\\talent_icons\\*.png , talent_bg\\*.png   icons and tree backgrounds from the client

Needs wow_common.py, wow_spelldata.py next to it, the DBC files from extract_dbc_multi.py
(Talent, TalentTab, SpellIcon ...), and  py -m pip install pillow mpyq pymysql openpyxl
Run:  py wow_talents.py
"""
import json
import os
import traceback

from wow_common import *   # noqa
from wow_spelldata import SpellData

XLSX = os.path.join(OUT_FOLDER, "wow_talents.xlsx")
JSON_DIR = os.path.join(OUT_FOLDER, "gamedata")
ICON_OUT = os.path.join(JSON_DIR, "talent_icons")
BG_OUT = os.path.join(JSON_DIR, "talent_bg")
CLASS_BY_MASK = dict(CLASSES)
CLASS_ORDER = ["Druid", "Hunter", "Mage", "Paladin", "Priest", "Rogue", "Shaman", "Warlock", "Warrior"]


def _mk(p):
    try:
        os.makedirs(p)
    except OSError:
        pass


def detect_prereq_layout(tal, ids, tab_of, row_of):
    """Talent.dbc: ranks at 4..12, then k prerequisite talents followed by k prerequisite ranks.
    Try k = 3, 2, 1 and keep the layout whose prerequisites look valid."""
    best, best_score = None, -1
    for k in (3, 2, 1):
        if 13 + 2 * k > tal.fc:
            continue
        score = bad = 0
        for r in tal.rows:
            for j in range(k):
                t, rk = r[13 + j], r[13 + k + j]
                if not t:
                    continue
                if t in ids and tab_of[t] == r[1] and row_of[t] < r[2] and rk < 9:
                    score += 1
                else:
                    bad += 1
        score -= bad * 3
        if score > best_score:
            best, best_score = k, score
    return best


def tab_fields(tab):
    def name(r):
        for i in range(1, 10):
            t = tab.s(r[i])
            if t:
                return t
        return ""
    cm = None
    for i in range(10, tab.fc):
        vals = [r[i] for r in tab.rows]
        if all(v in (1, 2, 4, 8, 16, 32, 64, 128, 256, 512, 1024) for v in vals) and len(set(vals)) >= 8:
            cm = i
            break
    bg = None
    for i in range(10, tab.fc):
        if any("talentframe" in tab.s(r[i]).lower() for r in tab.rows):
            bg = i
            break
    icon = 10 if tab.fc > 10 else None
    order = cm + 1 if cm is not None and cm + 1 < tab.fc else None
    return name, cm, bg, icon, order


def icon_names():
    d = load_dbc("SpellIcon")
    if not d:
        return {}
    out = {}
    for r in d.rows:
        t = d.s(r[1])
        if t:
            out[r[0]] = t.replace("/", "\\").split("\\")[-1].lower()
    return out


def png_from_blp(data, path, size=None):
    from PIL import Image
    import io
    im = Image.open(io.BytesIO(data)).convert("RGBA")
    if size:
        im = im.resize((size, size), Image.LANCZOS)
    im.save(path)


def open_client():
    try:
        import extract_maps
        arcs = extract_maps.open_archives()
        return lambda name: extract_maps.read_from(arcs, name)
    except Exception:
        traceback.print_exc()
        return None


def fetch_icons(reader, names):
    _mk(ICON_OUT)
    ok = miss = 0
    for n in sorted(names):
        p = os.path.join(ICON_OUT, n + ".png")
        if os.path.exists(p):
            ok += 1
            continue
        data = reader("Interface\\Icons\\%s.blp" % n) if reader else None
        if data:
            try:
                png_from_blp(data, p, 64)
                ok += 1
                continue
            except Exception:
                pass
        miss += 1
    log("  icons: %d ok, %d missing" % (ok, miss))


def fetch_background(reader, bgfile, key):
    """tree background = four 'TopLeft/TopRight/BottomLeft/BottomRight' textures"""
    if not reader or not bgfile:
        return ""
    from PIL import Image
    import io
    _mk(BG_OUT)
    path = os.path.join(BG_OUT, key + ".png")
    if os.path.exists(path):
        return "talent_bg/%s.png" % key
    parts = {}
    for suf in ("TopLeft", "TopRight", "BottomLeft", "BottomRight"):
        data = reader("%s-%s.blp" % (bgfile.replace("/", "\\"), suf))
        if not data:
            return ""
        parts[suf] = Image.open(io.BytesIO(data)).convert("RGB")
    w = parts["TopLeft"].width + parts["TopRight"].width
    h = parts["TopLeft"].height + parts["BottomLeft"].height
    im = Image.new("RGB", (w, h))
    im.paste(parts["TopLeft"], (0, 0))
    im.paste(parts["TopRight"], (parts["TopLeft"].width, 0))
    im.paste(parts["BottomLeft"], (0, parts["TopLeft"].height))
    im.paste(parts["BottomRight"], (parts["TopLeft"].width, parts["TopLeft"].height))
    im.save(path)
    return "talent_bg/%s.png" % key


def build(cur, sd, wb=None, save_workbook=True):
    tab, tal = load_dbc("TalentTab"), load_dbc("Talent")
    if not (tab and tal):
        raise RuntimeError("Talent.dbc / TalentTab.dbc missing - run extract_dbc_multi.py first")
    name_of, cm, bg, icon_f, order_f = tab_fields(tab)
    if cm is None:
        raise RuntimeError("could not find the class column in TalentTab.dbc")
    trees = {}
    for r in tab.rows:
        trees[r[0]] = {"id": r[0], "name": name_of(r), "class": CLASS_BY_MASK.get(r[cm], "?"),
                       "order": r[order_f] if order_f else 0,
                       "icon_id": r[icon_f] if icon_f else 0,
                       "background": tab.s(r[bg]) if bg else ""}
    ids = {r[0] for r in tal.rows}
    tab_of = {r[0]: r[1] for r in tal.rows}
    row_of = {r[0]: r[2] for r in tal.rows}
    k = detect_prereq_layout(tal, ids, tab_of, row_of)
    log("  Talent.dbc: %d fields, %d prerequisite slot(s) per talent" % (tal.fc, k or 0))
    talents = {}
    spell_ids = set()
    for r in tal.rows:
        if r[1] not in trees:
            continue
        ranks = []
        for s in r[4:13]:
            if not s:
                break
            ranks.append(s)
        pre = []
        if k:
            for j in range(k):
                if r[13 + j]:
                    pre.append({"talent": r[13 + j], "rank": r[13 + k + j] + 1})
        talents[r[0]] = {"id": r[0], "tree": r[1], "row": r[2], "col": r[3], "ranks": ranks, "prereq": pre}
        spell_ids.update(ranks)
    info = {}
    for rr in sd.fetch(cur, spell_ids):
        d = sd.row(rr)
        info[d["id"]] = d
    icons = icon_names()
    for t in talents.values():
        first = info.get(t["ranks"][0], {}) if t["ranks"] else {}
        t["name"] = first.get("name", "")
        t["icon"] = icons.get(first.get("icon_id"), "")
        t["rank_info"] = [{"spell": s, "text": info.get(s, {}).get("desc", ""),
                           "cost": info.get(s, {}).get("cost", ""), "cast": info.get(s, {}).get("cast", ""),
                           "cooldown": info.get(s, {}).get("cooldown", ""),
                           "range": info.get(s, {}).get("range", "")} for s in t["ranks"]]
    for tr in trees.values():
        tr["icon"] = icons.get(tr["icon_id"], "")
        tr["points"] = sum(len(t["ranks"]) for t in talents.values() if t["tree"] == tr["id"])
        tr["rows"] = 1 + max([t["row"] for t in talents.values() if t["tree"] == tr["id"]] or [0])
    # art from the game files
    reader = open_client()
    if reader:
        names = {t["icon"] for t in talents.values() if t["icon"]} | {tr["icon"] for tr in trees.values() if tr["icon"]}
        fetch_icons(reader, names)
        for tr in trees.values():
            try:
                tr["bg_file"] = fetch_background(reader, tr["background"], str(tr["id"]))
            except Exception:
                tr["bg_file"] = ""
    ordered = sorted(trees.values(), key=lambda t: (CLASS_ORDER.index(t["class"]) if t["class"] in CLASS_ORDER else 99, t["order"], t["id"]))

    # ------------------------------------------------------------ Excel
    if wb is None:
        from openpyxl import Workbook
        wb = Workbook()
        wb.remove(wb.active)
    for nm in ("Talent Trees", "Talents", "Talent Layout"):
        if nm in wb.sheetnames:
            del wb[nm]
    add_plain_sheet(wb, "Talent Trees", ["Class", "Tree", "Tree ID", "Order", "Talents", "Total Points", "Rows", "Icon"],
                    [11, 22, 8, 7, 8, 11, 6, 30],
                    [[t["class"], t["name"], t["id"], t["order"] + 1,
                      len([1 for x in talents.values() if x["tree"] == t["id"]]), t["points"], t["rows"], t["icon"]]
                     for t in ordered])
    rows = []
    for t in ordered:
        for x in sorted((x for x in talents.values() if x["tree"] == t["id"]), key=lambda x: (x["row"], x["col"])):
            pre = "; ".join("%s (%d points)" % (talents[p["talent"]]["name"], p["rank"]) for p in x["prereq"] if p["talent"] in talents)
            ri = x["rank_info"]
            txt = [ri[i]["text"] if i < len(ri) else "" for i in range(5)]
            rows.append([t["class"], t["name"], x["row"] + 1, x["col"] + 1, "%d points in tree" % (x["row"] * 5),
                         x["name"], len(x["ranks"]), pre, ri[0]["cost"] if ri else "", ri[0]["cast"] if ri else "",
                         ri[0]["cooldown"] if ri else "", ri[0]["range"] if ri else ""] + txt +
                        [", ".join(str(s) for s in x["ranks"]), x["icon"], x["id"]])
    add_plain_sheet(wb, "Talents",
                    ["Class", "Tree", "Row", "Column", "Needs", "Talent", "Max Rank", "Requires Talent", "Cost",
                     "Cast", "Cooldown", "Range", "Rank 1", "Rank 2", "Rank 3", "Rank 4", "Rank 5", "Rank Spell IDs",
                     "Icon", "Talent ID"],
                    [11, 20, 5, 7, 16, 28, 8, 30, 14, 8, 9, 9, 52, 52, 52, 52, 52, 22, 28, 8], rows)
    # grid layout like the game: 4 columns x N rows per tree
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter as gcl
    ws = wb.create_sheet("Talent Layout")
    ws.column_dimensions["A"].width = 14
    for c in range(2, 6):
        ws.column_dimensions[gcl(c)].width = 30
    rown = 1
    for t in ordered:
        ws.cell(row=rown, column=1, value="%s - %s" % (t["class"], t["name"])).font = Font(bold=True, color="FFFFFF")
        for c in range(1, 6):
            ws.cell(row=rown, column=c).fill = PatternFill("solid", fgColor="1F3864")
        rown += 1
        grid = {(x["row"], x["col"]): x for x in talents.values() if x["tree"] == t["id"]}
        for rr in range(t["rows"]):
            ws.cell(row=rown, column=1, value="Row %d (%d pts)" % (rr + 1, rr * 5)).font = Font(italic=True)
            for cc in range(4):
                x = grid.get((rr, cc))
                if x:
                    cell = ws.cell(row=rown, column=2 + cc, value="%s [0/%d]" % (x["name"], len(x["ranks"])))
                    if x["prereq"]:
                        cell.value += "\n(needs " + ", ".join(talents[p["talent"]]["name"] for p in x["prereq"] if p["talent"] in talents) + ")"
                    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
                    cell.fill = PatternFill("solid", fgColor="E9EFF7")
            ws.row_dimensions[rown].height = 34
            rown += 1
        rown += 1
    if save_workbook:
        wb.save(XLSX)

    # ------------------------------------------------------------ JSON + HTML
    _mk(JSON_DIR)
    data = {"trees": ordered, "talents": sorted(talents.values(), key=lambda x: (x["tree"], x["row"], x["col"]))}
    with open(os.path.join(JSON_DIR, "talents.json"), "w") as f:
        json.dump(data, f, indent=0)
    write_html(os.path.join(JSON_DIR, "talent_trees.html"), data)
    log("  talents: %d in %d trees  (%s)" % (len(talents), len(trees), JSON_DIR))
    return data


HTML = r"""<!doctype html><html><head><meta charset="utf-8"><title>WoW Classic talent trees</title>
<style>
body{background:#10131a;color:#ddd;font-family:Segoe UI,Arial,sans-serif;margin:0;padding:12px}
h1{font-size:18px;margin:0 0 8px}
.tabs button{background:#222a3a;color:#ccc;border:1px solid #3a4560;padding:6px 12px;margin:0 4px 4px 0;cursor:pointer;border-radius:3px}
.tabs button.on{background:#f0c040;color:#000}
.wrap{display:flex;flex-wrap:wrap;gap:14px;margin-top:10px}
.tree{width:316px;border:1px solid #3a4560;border-radius:6px;background:#161b27;padding:6px}
.tree h2{font-size:14px;margin:2px 4px 6px;display:flex;align-items:center;gap:6px}
.tree h2 img{width:24px;height:24px;border-radius:4px}
.tree h2 span{margin-left:auto;color:#f0c040;font-weight:bold}
.grid{position:relative;background-size:cover;border-radius:4px}
.grid svg{position:absolute;left:0;top:0;pointer-events:none}
.cell{position:absolute;width:48px;height:48px;border:2px solid #555;border-radius:6px;background:#000 center/cover;cursor:pointer;filter:grayscale(1) brightness(.6)}
.cell.avail{filter:none;border-color:#4c4}
.cell.full{filter:none;border-color:#f0c040}
.cell b{position:absolute;right:-4px;bottom:-6px;background:#000;border:1px solid #888;border-radius:8px;font-size:11px;padding:0 4px;color:#fff}
.cell.noicon{display:flex;align-items:center;justify-content:center;font-size:9px;text-align:center;color:#aaa;background:#222}
#tip{position:fixed;display:none;max-width:320px;background:#0a0c12ee;border:1px solid #888;border-radius:4px;padding:8px;font-size:12px;pointer-events:none;z-index:9}
#tip h3{margin:0 0 4px;font-size:14px;color:#f0c040}.dim{color:#8a8}.req{color:#e55}
.sum{margin:6px 0;color:#9ab}
</style></head><body><h1>WoW Classic talent trees &mdash; left click: add point, right click: remove</h1>
<div class="tabs" id="tabs"></div><div class="sum" id="sum"></div><div class="wrap" id="wrap"></div><div id="tip"></div>
<script>
const D=__DATA__;const T={};D.talents.forEach(t=>T[t.id]=t);
const pts={};D.talents.forEach(t=>pts[t.id]=0);
const classes=[...new Set(D.trees.map(t=>t.class))];let cur=classes[0];
const SX=62,SY=60,PX=14,PY=10;
function treePts(id){return D.talents.filter(t=>t.tree==id).reduce((a,t)=>a+pts[t.id],0)}
function allowed(t){const tp=treePts(t.tree);if(tp<t.row*5)return false;
 return t.prereq.every(p=>T[p.talent]&&pts[p.talent]>=p.rank)}
function valid(tree){const ts=D.talents.filter(t=>t.tree==tree);
 for(const t of ts){if(pts[t.id]>0){let before=0;ts.forEach(o=>{if(o.row<t.row)before+=pts[o.id]});
  if(before<t.row*5)return false;if(!t.prereq.every(p=>pts[p.talent]>=p.rank))return false}}return true}
function tabs(){const el=document.getElementById('tabs');el.innerHTML='';classes.forEach(c=>{const b=document.createElement('button');
 b.textContent=c;b.className=c==cur?'on':'';b.onclick=()=>{cur=c;render()};el.appendChild(b)});
 const r=document.createElement('button');r.textContent='Reset class';r.onclick=()=>{D.trees.filter(t=>t.class==cur).forEach(tr=>D.talents.filter(t=>t.tree==tr.id).forEach(t=>pts[t.id]=0));render()};el.appendChild(r)}
function tip(e,t){const el=document.getElementById('tip');if(!t){el.style.display='none';return}
 const n=pts[t.id],mx=t.ranks.length;let h='<h3>'+t.name+'</h3><div>Rank '+n+'/'+mx+'</div>';
 const ri=t.rank_info;const r0=ri[0]||{};
 if(r0.cost)h+='<div>'+r0.cost+'</div>';if(r0.range&&r0.range!='0')h+='<div>'+r0.range+' yd range</div>';
 if(r0.cast)h+='<div>'+r0.cast+(r0.cooldown?' cast, '+r0.cooldown+' cooldown':' cast')+'</div>';
 t.prereq.forEach(p=>{const ok=pts[p.talent]>=p.rank;h+='<div class="'+(ok?'dim':'req')+'">Requires '+p.rank+' points in '+(T[p.talent]||{}).name+'</div>'});
 if(t.row>0)h+='<div class="'+(treePts(t.tree)>=t.row*5?'dim':'req')+'">Requires '+t.row*5+' points in this tree</div>';
 if(n>0&&ri[n-1])h+='<p>'+ri[n-1].text+'</p>';else if(ri[0])h+='<p>'+ri[0].text+'</p>';
 if(n>0&&n<mx&&ri[n])h+='<p class="dim"><b>Next rank:</b> '+ri[n].text+'</p>';
 el.innerHTML=h;el.style.display='block';el.style.left=Math.min(e.clientX+16,innerWidth-340)+'px';el.style.top=Math.min(e.clientY+16,innerHeight-200)+'px'}
function render(){tabs();const w=document.getElementById('wrap');w.innerHTML='';let total=0;
 D.trees.filter(t=>t.class==cur).forEach(tr=>{const tp=treePts(tr.id);total+=tp;
  const box=document.createElement('div');box.className='tree';
  box.innerHTML='<h2>'+(tr.icon?'<img src="talent_icons/'+tr.icon+'.png">':'')+tr.name+'<span>'+tp+'</span></h2>';
  const g=document.createElement('div');g.className='grid';g.style.width=(PX*2+SX*3+48)+'px';g.style.height=(PY*2+SY*(tr.rows-1)+48)+'px';
  if(tr.bg_file)g.style.backgroundImage='url('+tr.bg_file+')';
  let svg='<svg width="'+g.style.width+'" height="'+g.style.height+'"><defs><marker id="a" markerWidth="6" markerHeight="6" refX="3" refY="3" orient="auto"><path d="M0,0 L6,3 L0,6 z" fill="#eee"/></marker></defs>';
  D.talents.filter(t=>t.tree==tr.id).forEach(t=>t.prereq.forEach(p=>{const s=T[p.talent];if(!s)return;
   const x1=PX+s.col*SX+24,y1=PY+s.row*SY+24,x2=PX+t.col*SX+24,y2=PY+t.row*SY+24;
   svg+='<line x1="'+x1+'" y1="'+y1+'" x2="'+x2+'" y2="'+y2+'" stroke="#eee" stroke-width="3" opacity=".7"/>'}));
  g.innerHTML=svg+'</svg>';
  D.talents.filter(t=>t.tree==tr.id).forEach(t=>{const c=document.createElement('div');
   const n=pts[t.id],mx=t.ranks.length;c.className='cell'+(n==mx?' full':(n>0||allowed(t)?' avail':''));
   if(!t.icon)c.className+=' noicon';c.style.left=(PX+t.col*SX)+'px';c.style.top=(PY+t.row*SY)+'px';
   if(t.icon)c.style.backgroundImage='url(talent_icons/'+t.icon+'.png)';else c.textContent=t.name;
   c.innerHTML+='<b>'+n+'/'+mx+'</b>';
   c.onmouseenter=c.onmousemove=e=>tip(e,t);c.onmouseleave=()=>tip(null,null);
   c.onclick=()=>{if(n<mx&&allowed(t)&&total<51){pts[t.id]++;render()}};
   c.oncontextmenu=e=>{e.preventDefault();if(n>0){pts[t.id]--;if(!valid(tr.id))pts[t.id]++;render()}};
   g.appendChild(c)});
  box.appendChild(g);w.appendChild(box)});
 document.getElementById('sum').textContent=cur+': '+total+' points spent - required level '+(total?total+9:1)}
render();
</script></body></html>"""


def write_html(path, data):
    with open(path, "w", encoding="utf-8") as f:
        f.write(HTML.replace("__DATA__", json.dumps(data)))


def main():
    log("Connecting to database ...")
    conn = connect()
    cur = conn.cursor()
    sd = SpellData(cur)
    build(cur, sd)
    conn.close()
    log("Done.  Open  %s  in your browser for the clickable trees." % os.path.join(JSON_DIR, "talent_trees.html"))
    log("Excel: %s" % XLSX)


if __name__ == "__main__":
    run(main)
