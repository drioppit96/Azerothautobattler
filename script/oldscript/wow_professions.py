# -*- coding: utf-8 -*-
"""
Profession spreadsheets:
  C:\\wow excel\\wow_profession_items.xlsx   (Profession Items, Sources)
  C:\\wow excel\\wow_recipes.xlsx            (Recipes, Materials)

Needs wow_common.py in the same folder, the SPP Database window open, and the
DBC files from extract_dbc_multi.py.   Run:  py wow_professions.py
"""
import os
from collections import defaultdict

from wow_common import *   # noqa

OUT_ITEMS = os.path.join(OUT_FOLDER, "wow_profession_items.xlsx")
OUT_RECIPES = os.path.join(OUT_FOLDER, "wow_recipes.xlsx")

PROFESSIONS = {171: "Alchemy", 164: "Blacksmithing", 333: "Enchanting",
               202: "Engineering", 165: "Leatherworking", 197: "Tailoring",
               185: "Cooking", 129: "First Aid", 186: "Mining", 40: "Poisons"}
EFFECT_CREATE_ITEM, EFFECT_ENCHANT_PERM, EFFECT_ENCHANT_TEMP, EFFECT_LEARN = 24, 53, 54, 36
FOCUS = {1: "Anvil", 2: "Loom", 3: "Forge"}


def main():
    log("Connecting to database ...")
    conn = connect()
    cur = conn.cursor()

    # ---- skill names and the recipe -> skill table from the client
    names = dict(PROFESSIONS)
    sk = load_dbc("SkillLine")
    if sk:
        idx = sk.find_field(171, "Alchemy")
        if idx is not None:
            for r in sk.rows:
                if r[0] in PROFESSIONS:
                    names[r[0]] = sk.s(r[idx]) or PROFESSIONS[r[0]]
    sla = load_dbc("SkillLineAbility")
    if not sla:
        log("SkillLineAbility.dbc is missing - run extract_dbc_multi.py first.")
        return
    if sla.fc != 14:
        log("  note: SkillLineAbility has %d fields (expected 14) - results may be off" % sla.fc)
    abil = {}
    for r in sla.rows:
        skill, spell = r[1], r[2]
        if skill in PROFESSIONS and spell not in abil:
            abil[spell] = {"skill": skill, "min": r[7], "high": r[10], "low": r[11]}
    log("Profession abilities in client data: %d" % len(abil))

    # ---- spell data from the database
    sc = columns(cur, "spell_template")
    spells = {}
    for part in chunks(list(abil)):
        cur.execute("SELECT * FROM spell_template WHERE `%s` IN (%s)" % (sc["id"], ph(len(part))), part)
        for r in cur.fetchall():
            spells[g(r, sc, "id")] = r

    def S(r, name, d=0):
        return g(r, sc, name, d)

    recipes = {}
    for sid, r in spells.items():
        for e in (1, 2, 3):
            eff = S(r, "effect%d" % e)
            if eff == EFFECT_CREATE_ITEM and S(r, "effectitemtype%d" % e):
                base, dice = S(r, "effectbasepoints%d" % e), S(r, "effectbasedice%d" % e)
                qty = max(1, base + max(dice, 1)) if base >= 0 else 1
                recipes[sid] = {"item": S(r, "effectitemtype%d" % e), "qty": qty, "kind": "Item"}
                break
            if eff in (EFFECT_ENCHANT_PERM, EFFECT_ENCHANT_TEMP):
                recipes[sid] = {"item": 0, "qty": 1, "kind": "Enchantment"}
                break
    log("Craftable recipes found: %d" % len(recipes))

    # ---- who teaches them: learn-spells, recipe items, trainers
    learn_to_recipe = {}
    cur.execute("SELECT * FROM spell_template WHERE Effect1 = %s OR Effect2 = %s OR Effect3 = %s",
                (EFFECT_LEARN, EFFECT_LEARN, EFFECT_LEARN))
    for r in cur.fetchall():
        for e in (1, 2, 3):
            if S(r, "effect%d" % e) == EFFECT_LEARN and S(r, "effecttriggerspell%d" % e) in recipes:
                learn_to_recipe[S(r, "id")] = S(r, "effecttriggerspell%d" % e)

    ic = columns(cur, "item_template")
    cur.execute("SELECT * FROM item_template WHERE `%s` = 9" % ic["class"])
    recipe_item_rows = cur.fetchall()
    recipe_item = {}      # recipe spell -> item id
    for r in recipe_item_rows:
        for n in range(1, 6):
            sid = g(r, ic, "spellid_%d" % n)
            if sid in learn_to_recipe:
                recipe_item.setdefault(learn_to_recipe[sid], g(r, ic, "entry"))
            elif sid in recipes:
                recipe_item.setdefault(sid, g(r, ic, "entry"))

    teach_ids = set(recipes) | set(learn_to_recipe)
    trainers = defaultdict(list)   # recipe spell -> [(name, cost, req skill, req level)]
    ct = columns(cur, "creature_template")
    tnames = {}
    tpl_col = ct.get("trainertemplateid")
    cur.execute("SELECT `%s` e, `%s` n, %s t FROM creature_template"
                % (ct["entry"], ct["name"], ("`%s`" % tpl_col) if tpl_col else "0"))
    by_tpl = defaultdict(list)
    for r in cur.fetchall():
        tnames[r["e"]] = r["n"]
        if r["t"]:
            by_tpl[r["t"]].append(r["n"])
    for tbl in ("npc_trainer", "npc_trainer_template"):
        if not table_exists(cur, tbl):
            continue
        tc = columns(cur, tbl)
        for part in chunks(list(teach_ids)):
            cur.execute("SELECT * FROM `%s` WHERE `%s` IN (%s)" % (tbl, tc["spell"], ph(len(part))), part)
            for r in cur.fetchall():
                spell = g(r, tc, "spell")
                rid = learn_to_recipe.get(spell, spell)
                if rid not in recipes:
                    continue
                ent = g(r, tc, "entry")
                if tbl == "npc_trainer":
                    nms = [tnames.get(ent, "NPC #%d" % ent)]
                else:
                    nms = by_tpl.get(ent, ["Trainer template #%d" % ent])
                for nm in nms:
                    trainers[rid].append((nm, g(r, tc, "spellcost"), g(r, tc, "reqskillvalue"),
                                          g(r, tc, "reqlevel")))

    # ---- items involved
    reagents = defaultdict(list)        # item -> [(recipe spell, qty)]
    tools = {}
    for sid in recipes:
        r = spells[sid]
        for n in range(1, 9):
            if S(r, "reagent%d" % n) > 0:
                reagents[S(r, "reagent%d" % n)].append((sid, S(r, "reagentcount%d" % n) or 1))
        for n in (1, 2):
            if S(r, "totem%d" % n):
                tools.setdefault(sid, []).append(S(r, "totem%d" % n))
    need = set(reagents) | {v["item"] for v in recipes.values() if v["item"]}
    need |= set(recipe_item.values())
    for lst in tools.values():
        need.update(lst)
    cur.execute("SELECT `%s` e FROM item_template WHERE `%s` IN (7, 9)" % (ic["entry"], ic["class"]))
    need |= {r["e"] for r in cur.fetchall()}
    items = load_items_by_ids(cur, need)
    log("Profession-related items: %d" % len(items))

    # per-item relations
    created_by = defaultdict(list)
    for sid, v in recipes.items():
        if v["item"]:
            created_by[v["item"]].append(sid)

    def spell_name(sid):
        r = spells.get(sid)
        return S(r, "spellname", "") if r else "spell %d" % sid

    def prof_of(sid):
        return names.get(abil[sid]["skill"], "")

    cls9_prof = {0: "", 1: "Leatherworking", 2: "Tailoring", 3: "Engineering",
                 4: "Blacksmithing", 5: "Cooking", 6: "Alchemy", 7: "First Aid",
                 8: "Enchanting", 9: "Fishing"}

    def item_professions(it):
        p = set()
        if it["cls"] == 9 and cls9_prof.get(it["sub"]):
            p.add(cls9_prof[it["sub"]])
        for sid in created_by.get(it["id"], []):
            p.add(prof_of(sid))
        for sid, _ in reagents.get(it["id"], []):
            p.add(prof_of(sid))
        return ", ".join(sorted(p))

    log("Looking up item sources (drops, vendors, quests) ...")
    sources = build_sources(cur, items)
    for iid, lst in created_by.items():
        if iid in items:
            for sid in lst:
                sources.append((iid, "Crafted", "%s: %s" % (prof_of(sid), spell_name(sid)),
                                "skill %d" % abil[sid]["min"]))
    by_item = defaultdict(list)
    for it, typ, nme, detail in sources:
        by_item[it].append((typ, nme, detail))
    conn.close()

    from openpyxl import Workbook
    icons = Icons()

    # =========================== profession items workbook
    wb = Workbook()
    ordered = sorted(items.values(), key=lambda x: (x["cls"], x["sub"], x["reqlvl"], x["name"]))

    def extra(it):
        cb = created_by.get(it["id"], [])
        used = reagents.get(it["id"], [])
        crafted = "; ".join("%s: %s" % (prof_of(s), spell_name(s)) for s in cb[:3])
        if len(cb) > 3:
            crafted += " (+%d more)" % (len(cb) - 3)
        usedin = "; ".join(spell_name(s) for s, _ in used[:4])
        if len(used) > 4:
            usedin += " (+%d more)" % (len(used) - 4)
        srcs = by_item.get(it["id"], [])
        detail = "; ".join("%s: %s (%s)" % (t, n_, d) for t, n_, d in srcs[:5])
        if len(srcs) > 5:
            detail += " (+%d more)" % (len(srcs) - 5)
        return [item_professions(it), crafted, len(used) or "", usedin,
                source_summary(srcs) or "unknown", detail]

    write_item_sheet(wb, "Profession Items", ordered, icons,
                     extra_headers=["Profession(s)", "Crafted By", "Used In (recipes)",
                                    "Used In (names)", "Sources (summary)", "Sources (detail)"],
                     extra_widths=[26, 44, 10, 50, 30, 80], extra_fn=extra, first=True)
    rows = [[it, items[it]["name"], t, n_, d] for it, t, n_, d in
            sorted(sources, key=lambda s: (s[0], s[1])) if it in items]
    add_plain_sheet(wb, "Sources", ["Item ID", "Item Name", "Source Type", "Source", "Details"],
                    [10, 36, 14, 60, 40], rows)
    wb.save(OUT_ITEMS)
    log("Saved %s" % OUT_ITEMS)

    # =========================== recipes workbook
    wb2 = Workbook()
    rec_cols = ["Icon", "Profession", "Recipe", "Spell ID", "Creates", "Item ID", "Qty",
                "Quality", "Skill: Orange", "Yellow", "Green", "Grey", "Materials",
                "Tools", "Required Station", "Learned From", "Trainer(s)",
                "Trainer Cost", "Trainer Req. Level", "Recipe Item", "Recipe Item Sources",
                "Item Link", "Spell Link"]
    widths = [8, 16, 38, 9, 34, 8, 6, 10, 10, 8, 8, 8, 60, 24, 12, 18, 40, 12, 10, 40, 50, 34, 34]
    ws = new_sheet(wb2, "Recipes", rec_cols, widths, first=True)
    mat_rows = []

    def sorted_key(sid):
        return (prof_of(sid), abil[sid]["min"], spell_name(sid))

    n = 1
    for sid in sorted(recipes, key=sorted_key):
        n += 1
        v, a, r = recipes[sid], abil[sid], spells[sid]
        it = items.get(v["item"]) if v["item"] else None
        mats = []
        for k in range(1, 9):
            rid = S(r, "reagent%d" % k)
            if rid > 0:
                cnt = S(r, "reagentcount%d" % k) or 1
                rn = items[rid]["name"] if rid in items else "item %d" % rid
                mats.append("%dx %s" % (cnt, rn))
                mat_rows.append((sid, rid, cnt))
        tool_txt = ", ".join(items[t]["name"] if t in items else "item %d" % t
                             for t in tools.get(sid, []))
        focus = FOCUS.get(S(r, "requiresspellfocus"), ("Station #%d" % S(r, "requiresspellfocus"))
                          if S(r, "requiresspellfocus") else "")
        tr = trainers.get(sid, [])
        ritem = recipe_item.get(sid)
        learned = []
        if tr:
            learned.append("Trainer")
        if ritem:
            learned.append("Recipe item")
        if not learned:
            learned.append("Starting / quest / unknown")
        tr_names = sorted({t[0] for t in tr})
        tr_txt = ", ".join(tr_names[:4]) + (" (+%d more)" % (len(tr_names) - 4) if len(tr_names) > 4 else "")
        cost = money(min(t[1] for t in tr)) if tr else ""
        reqlvl = min(t[3] for t in tr) if tr else ""
        ritem_txt = ritem_src = ""
        if ritem and ritem in items:
            ritem_txt = "%s (%d)" % (items[ritem]["name"], ritem)
            ritem_src = source_summary(by_item.get(ritem, [])) or "unknown"
            first = "; ".join("%s: %s" % (t, n_) for t, n_, _ in by_item.get(ritem, [])[:3])
            if first:
                ritem_src += " - " + first
        row = [
            "", prof_of(sid), spell_name(sid) + ((" (%s)" % S(r, "rank1", "")) if S(r, "rank1", "") else ""),
            sid, (it["name"] if it else ("Enchantment" if v["kind"] == "Enchantment" else "")),
            v["item"] or "", v["qty"] if v["item"] else "",
            QUALITY.get(it["quality"], "") if it else "",
            a["min"], (a["low"] + a["high"]) // 2 if a["high"] else "", a["low"], a["high"],
            ", ".join(mats), tool_txt, focus, ", ".join(learned), tr_txt, cost, reqlvl,
            ritem_txt, ritem_src,
            ("https://www.wowhead.com/classic/item=%d" % v["item"]) if v["item"] else "",
            "https://www.wowhead.com/classic/spell=%d" % sid]
        ws.append(row)
        icon = icons.for_item(it["display"]) if it else icons.for_spell(S(r, "spelliconid"))
        style_row(ws, n, len(rec_cols), icon_path=icon, tall=True)
        if it:
            from openpyxl.styles import Font
            ws.cell(row=n, column=5).font = Font(bold=True, color=QUALITY_COLOR.get(it["quality"], "000000"))
        if n % 100 == 0:
            log("  recipes written: %d / %d" % (n - 1, len(recipes)))
    finish_sheet(ws, len(rec_cols), freeze="D2")

    # materials sheet (one row per reagent, with icon)
    from openpyxl.styles import Font
    mcols = ["Icon", "Profession", "Recipe", "Spell ID", "Creates", "Material", "Material ID",
             "Qty", "Material Quality", "Material Sources", "Material Link"]
    mws = new_sheet(wb2, "Materials", mcols, [8, 16, 38, 9, 30, 32, 10, 6, 12, 50, 34])
    m = 1
    for sid, rid, cnt in sorted(mat_rows, key=lambda t: (sorted_key(t[0]), t[1])):
        m += 1
        it = items.get(rid)
        cr = items.get(recipes[sid]["item"]) if recipes[sid]["item"] else None
        link = "https://www.wowhead.com/classic/item=%d" % rid
        mws.append(["", prof_of(sid), spell_name(sid), sid, cr["name"] if cr else "",
                    it["name"] if it else "item %d" % rid, rid, cnt,
                    QUALITY.get(it["quality"], "") if it else "",
                    source_summary(by_item.get(rid, [])) or "", link])
        style_row(mws, m, len(mcols), link_col=len(mcols), link=link,
                  icon_path=icons.for_item(it["display"]) if it else None, tall=True)
    finish_sheet(mws, len(mcols), freeze="D2")
    wb2.save(OUT_RECIPES)
    if icons.missing:
        log("Icons missing for %d rows." % icons.missing)
    log("Saved %s" % OUT_RECIPES)
    log("Done!")


if __name__ == "__main__":
    run(main)
